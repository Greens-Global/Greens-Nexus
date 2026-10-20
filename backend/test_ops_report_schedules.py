"""Operations > Reports > Scheduled (ops_report_schedules.py and the
/ops-reports/schedules endpoints). Graph is mocked; Cubby is faked; the rows
are test_ops_reports' seed. Run alone: python -m pytest test_ops_report_schedules.py"""
import io
from datetime import date, datetime
from types import SimpleNamespace

import openpyxl
import pytest

import graph_mail
import models
import ops_report_builder as builder
import ops_report_schedules as sched
import ops_reports
from database import SessionLocal
from test_ops_reports import SECRETS, TODAY, FakeCubby, _client, _seed

TZ = sched.ZoneInfo(sched.TZ)
OPS = "ops.manager@greensglobal.com"
OTHER = "other.manager@greensglobal.com"


@pytest.fixture
def client(monkeypatch):
    import cubby
    _seed()
    db = SessionLocal()
    try:
        db.query(models.OpsReportSchedule).delete()
        db.commit()
    finally:
        db.close()
    monkeypatch.setenv("CUBBY_API_KEY", "test-key-not-real")
    monkeypatch.setattr(cubby, "_facility_today", lambda fac: TODAY)
    monkeypatch.setattr(ops_reports, "_cubby", FakeCubby())
    ops_reports._ledger_cache.clear()
    return _client()


@pytest.fixture
def mail(monkeypatch):
    sent = []
    monkeypatch.setattr(graph_mail, "graph_configured", lambda: True)
    monkeypatch.setattr(graph_mail, "DEFAULT_FROM_EMAIL", "nexus@greensglobal.com")
    monkeypatch.setattr(graph_mail, "send_mail", lambda **kw: sent.append(kw) or {})
    return sent


def _body(**kw):
    b = {"name": "Weekly Rent Roll", "kind": "rent-roll", "facilityIds": ["fac_a", "fac_b"],
         "period": {"mode": "previous", "unit": "week"}, "format": "xlsx", "frequency": "weekly", "day": 0, "hour": 7,
         "recipients": [OPS, OTHER]}
    b.update(kw)
    return b


# ── Periods ──────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("mode,unit,today,expected", [
    ("current", "day", date(2026, 10, 20), ("2026-10-20", "2026-10-20")),
    ("previous", "day", date(2026, 10, 1), ("2026-09-30", "2026-09-30")),
    ("current", "week", date(2026, 10, 22), ("2026-10-19", "2026-10-22")),        # Monday start
    ("previous", "week", date(2026, 10, 19), ("2026-10-12", "2026-10-18")),
    ("current", "month", date(2026, 10, 20), ("2026-10-01", "2026-10-20")),
    ("previous", "month", date(2026, 3, 15), ("2026-02-01", "2026-02-28")),
    ("previous", "month", date(2026, 1, 5), ("2025-12-01", "2025-12-31")),
    ("current", "quarter", date(2026, 11, 2), ("2026-10-01", "2026-11-02")),
    ("previous", "quarter", date(2026, 2, 10), ("2025-10-01", "2025-12-31")),
    ("current", "year", date(2026, 10, 20), ("2026-01-01", "2026-10-20")),
    ("previous", "year", date(2026, 10, 20), ("2025-01-01", "2025-12-31")),
])
def test_resolve_period(mode, unit, today, expected):
    assert sched.resolve_period({"mode": mode, "unit": unit}, today) == expected


def test_resolve_period_refuses_unknown():
    with pytest.raises(sched.ScheduleError):
        sched.resolve_period({"mode": "next", "unit": "month"}, date(2026, 10, 20))


def _row(**kw):
    base = dict(frequency="daily", day=0, hour=7, include_weekends=False, active=True, last_sent_period="")
    base.update(kw)
    return SimpleNamespace(**base)


def test_due_period_and_next_run():
    sat_8 = datetime(2026, 10, 24, 8, 0, tzinfo=TZ)
    mon_6 = datetime(2026, 10, 19, 6, 0, tzinfo=TZ)
    mon_8 = datetime(2026, 10, 19, 8, 0, tzinfo=TZ)
    assert sched.due_period(_row(), mon_6) == ""
    assert sched.due_period(_row(), mon_8) == "2026-10-19"
    assert sched.due_period(_row(), sat_8) == ""                                   # weekdays only
    assert sched.due_period(_row(include_weekends=True), sat_8) == "2026-10-24"
    assert sched.due_period(_row(frequency="weekly", day=2), mon_8) == ""           # Wednesday, not yet
    assert sched.due_period(_row(frequency="weekly", day=0), sat_8) == "2026-10-19"  # missed Monday: still owed
    assert sched.due_period(_row(frequency="monthly", day=1), mon_8) == "2026-10"
    # Next run: Saturday's daily weekday schedule goes Monday 7 AM Pacific.
    assert sched.next_run_at(_row(), sat_8) == "2026-10-26T07:00:00-07:00"
    assert sched.next_run_at(_row(last_sent_period="2026-10-19"), mon_8) == "2026-10-20T07:00:00-07:00"
    assert sched.next_run_at(_row(), mon_8).startswith("2026-10-19T08:00")          # owed now
    assert sched.next_run_at(_row(frequency="monthly", day=1, last_sent_period="2026-10"), mon_8) == "2026-11-01T07:00:00-08:00"   # DST ends that morning
    assert sched.next_run_at(_row(active=False), mon_8) == ""


# ── CRUD and permissions ─────────────────────────────────────────────────────

def test_create_list_update_delete_and_permissions(client):
    r = client.post("/ops-reports/schedules", json=_body())
    assert r.status_code == 201, r.text
    s = r.json()["schedule"]
    assert s["title"] == "Rent Roll" and s["facilityIds"] == ["fac_a", "fac_b"] and s["mine"] and s["canEdit"]
    assert s["period"] == {"mode": "previous", "unit": "week"} and s["format"] == "xlsx" and s["active"]
    assert s["recipients"] == [OPS, OTHER] and s["nextRunAt"] and s["describe"].startswith("Every Monday at 7:00 AM")
    sid = s["id"]
    # A supervisor (viewer) cannot create; an employee cannot even list.
    sup = _client(level=2, email="sup@greensglobal.com")
    assert sup.post("/ops-reports/schedules", json=_body()).status_code == 401
    assert sup.get("/ops-reports/schedules").json()["schedules"] == []
    assert _client(level=1, email="emp@greensglobal.com").get("/ops-reports/schedules").status_code == 401
    # A recipient sees it but cannot change it; a level-4 can.
    other = _client(email=OTHER)
    seen = other.get("/ops-reports/schedules").json()["schedules"]
    assert [x["id"] for x in seen] == [sid] and seen[0]["canEdit"] is False and seen[0]["mine"] is False
    assert other.patch(f"/ops-reports/schedules/{sid}", json={"name": "Hijack"}).status_code == 403
    boss = _client(level=4, email="boss@greensglobal.com")
    assert boss.patch(f"/ops-reports/schedules/{sid}", json={"active": False}).json()["schedule"]["nextRunAt"] == ""
    up = client.patch(f"/ops-reports/schedules/{sid}", json={"frequency": "monthly", "day": 1, "format": "pdf", "active": True})
    assert up.status_code == 200 and up.json()["schedule"]["format"] == "pdf"
    assert client.delete(f"/ops-reports/schedules/{sid}").status_code == 200
    assert client.get("/ops-reports/schedules").json()["schedules"] == []


@pytest.mark.parametrize("bad", [
    {"kind": "nope"}, {"kind": "inventory-activity"}, {"recipients": []}, {"recipients": ["someone@gmail.com"]},
    {"recipients": ["gone@greensglobal.com"]}, {"frequency": "hourly"}, {"frequency": "weekly", "day": 7},
    {"frequency": "monthly", "day": 31}, {"hour": 24}, {"format": "docx"}, {"period": {"mode": "next", "unit": "month"}},
    {"facilityIds": ["fac_zzz"]}, {"name": "  "}, {"options": {"x": {"nested": True}}},
])
def test_create_validates(client, bad):
    assert client.post("/ops-reports/schedules", json=_body(**bad)).status_code == 422


# ── Sending ──────────────────────────────────────────────────────────────────

def _make(client, **kw):
    sid = client.post("/ops-reports/schedules", json=_body(**kw)).json()["schedule"]["id"]
    db = SessionLocal()
    try:
        db.get(models.OpsReportSchedule, sid).last_sent_period = ""      # as if made before this period
        db.commit()
    finally:
        db.close()
    return sid


def test_scan_sends_once_per_period_with_the_file_attached_and_retries(client, mail, monkeypatch):
    sid = _make(client)
    now = datetime(2026, 10, 19, 8, 0, tzinfo=TZ)
    assert sched.scan_once(now) == 1
    assert sched.scan_once(now) == 0                                            # once per period
    m = mail[0]
    assert m["to"] == [OPS, OTHER]
    assert m["subject"] == "Ops Report: Weekly Rent Roll (As of 10/15/2026)"
    name, ctype, data = m["attachments"][0]
    assert name.startswith("Rent Roll - ") and name.endswith(".xlsx")
    assert ctype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    text = " ".join(str(ws.cell(i, c).value) for i in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1))
    assert "Aiden Robb" in text and "Cy Park" in text
    assert "attached" in m["html"] and "Total Monthly Rent" in m["html"]
    for s in SECRETS:
        assert s not in text and s not in m["html"]
    # Next week: a failure gives the claim back and waits an hour.
    def boom(**kw):
        raise graph_mail.GraphMailError("Graph is down")
    monkeypatch.setattr(graph_mail, "send_mail", boom)
    later = datetime(2026, 10, 26, 8, 0, tzinfo=TZ)
    assert sched.scan_once(later) == 0
    db = SessionLocal()
    try:
        row = db.get(models.OpsReportSchedule, sid)
        assert row.last_sent_period == "2026-10-19" and "Graph is down" in row.last_error
    finally:
        db.close()
    monkeypatch.setattr(graph_mail, "send_mail", lambda **kw: mail.append(kw) or {})
    assert sched.scan_once(later) == 0
    monkeypatch.setattr(sched, "RETRY_AFTER_SEC", -1)
    assert sched.scan_once(later) == 1
    assert client.get("/ops-reports/schedules").json()["schedules"][0]["lastError"] == ""


def test_timed_report_uses_the_resolved_period_and_pdf(client, mail):
    _make(client, kind="payment-detail", period={"mode": "current", "unit": "month"}, format="pdf", frequency="daily",
          facilityIds=["fac_a"])
    assert sched.scan_once(datetime(2026, 10, 15, 9, 0, tzinfo=TZ)) == 1
    m = mail[0]
    assert "(10/01/2026 through 10/15/2026)" in m["subject"]
    name, ctype, data = m["attachments"][0]
    assert ctype == "application/pdf" and data.startswith(b"%PDF") and name.endswith(".pdf")


def test_paused_schedules_and_new_ones_wait_for_the_next_occurrence(client, mail):
    sid = client.post("/ops-reports/schedules", json=_body(frequency="daily")).json()["schedule"]["id"]
    # Before 7 AM nothing is due; after it, today's send passed before the
    # schedule existed - either way nothing goes now.
    assert sched.scan_once(sched.local_now()) == 0 and not mail
    client.patch(f"/ops-reports/schedules/{sid}", json={"active": False})
    assert sched.scan_once(datetime(2030, 1, 2, 9, 0, tzinfo=TZ)) == 0                      # paused


def test_a_file_too_big_for_email_is_an_error(client, mail, monkeypatch):
    _make(client)
    monkeypatch.setattr(sched, "MAX_ATTACH_BYTES", 100)
    assert sched.scan_once(datetime(2026, 10, 19, 8, 0, tzinfo=TZ)) == 0
    err = client.get("/ops-reports/schedules").json()["schedules"][0]["lastError"]
    assert "more than email takes" in err and not mail


def test_send_now_to_me_and_to_everyone(client, mail):
    sid = _make(client, format="csv")
    r = client.post(f"/ops-reports/schedules/{sid}/send", json={"to_me": True})
    assert r.status_code == 200, r.text
    assert mail[-1]["to"] == [OPS] and mail[-1]["attachments"][0][1] == "text/csv"
    assert mail[-1]["attachments"][0][2].decode("utf-8-sig").startswith("Facility,Unit")
    assert client.get("/ops-reports/schedules").json()["schedules"][0]["lastSentAt"] == ""   # a test send is not the send
    assert client.post(f"/ops-reports/schedules/{sid}/send", json={}).json()["to"] == [OPS, OTHER]
    assert client.get("/ops-reports/schedules").json()["schedules"][0]["lastSentAt"]


def test_builder_loop_calls_the_schedule_scan():
    import inspect
    assert "ops_report_schedules.scan_once" in inspect.getsource(builder.ops_report_schedule_loop)
