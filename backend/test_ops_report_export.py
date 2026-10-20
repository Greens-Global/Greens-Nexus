"""Operations > Reports downloads (ops_report_export.py and the /export
endpoints). Pure renders on made-up results, then the endpoints on the rows
test_ops_reports seeds (Cubby faked). Run alone:
python -m pytest test_ops_report_export.py"""
import csv
import io
from datetime import date, datetime

import openpyxl
import pytest

import ops_report_export as export
import ops_reports
from test_ops_reports import SECRETS, _client, _seed, _spec, FakeCubby, TODAY  # noqa: F401 - shared seed

NOW = datetime(2026, 10, 20, 15, 5, tzinfo=export.ZoneInfo(export.TZ))


def col(k, label, t="text", total=False):
    return {"key": k, "label": label, "type": t, "total": total}


def multi_result():
    cols = [col("facility", "Facility"), col("tenant", "Tenant", "tenant"), col("unit", "Unit", "unit"),
            col("moveIn", "Move-In Date", "date"), col("rent", "Monthly Rent", "money", True), col("occ", "Occupancy", "pct"),
            col("count", "Units", "int", True)]
    rows = [{"_f": "fa", "_c": "c1", "facility": "GS Temecula", "tenant": "Ada Sample", "unit": "A1", "moveIn": "2025-03-01",
             "rent": 15000, "occ": 85.7, "count": 1},
            {"_f": "fa", "_c": "c2", "facility": "GS Temecula", "tenant": "=cmd|' /C calc'!A0", "unit": "A2", "moveIn": "2026-10-02",
             "rent": 0, "occ": None, "count": 1},
            {"_f": "fb", "_c": "c3", "facility": "GS Murrieta", "tenant": "Cy Park", "unit": "B1", "moveIn": "2026-01-10",
             "rent": 10000, "occ": 50.0, "count": 1},
            {"_f": "fb", "_c": "c4", "facility": "GS Murrieta", "tenant": "Dee Fox", "unit": "B2", "moveIn": "",
             "rent": -2500, "occ": 0, "count": 1}]
    return {"key": "rent-roll", "title": "Rent Roll", "period": {"from": "2026-10-01", "to": "2026-10-20"}, "asOf": "2026-10-20",
            "facility": "All Facilities", "facilities": ["GS Murrieta", "GS Temecula"], "columns": cols, "rows": rows,
            "totals": {"rent": 22500, "count": 4}, "groups": {}, "notes": ["A note about the rows."],
            "summary": [{"label": "Rent Roll", "value": 22500, "type": "money"}, {"label": "Occupancy", "value": 71.4, "type": "pct"}],
            "syncedAt": "2026-10-20T21:00:00+00:00", "available": True}


def wide_result(n_rows=120):
    cols = [col("facility", "Facility"), col("type", "Account Type"), col("tenant", "Tenant", "tenant"), col("unit", "Unit", "unit")]
    labels = ["Beginning AR Balance", "Rent", "Insurance", "Fees", "State Tax", "Local Tax", "Prepaid", "Deposit", "Other",
              "Waived Rent", "Waived Fees", "Discounts", "Payments", "Refunds", "Reversals", "Ending AR Balance"]
    cols += [col(f"m{i}", lab, "money", True) for i, lab in enumerate(labels)]
    rows = [{"facility": "GS Temecula", "type": "Individual", "tenant": "Christopher Montgomery-Wellington the Third " * 3,
             "unit": f"A{i}", **{f"m{j}": 1234567 * (j % 3) for j in range(len(labels))}} for i in range(n_rows)]
    return {"key": "accounts-receivable", "title": "Accounts Receivable", "period": {"from": "2026-10-01", "to": "2026-10-20"},
            "asOf": "2026-10-20", "facility": "GS Temecula", "columns": cols, "rows": rows,
            "totals": {f"m{j}": 1234567 * (j % 3) * n_rows for j in range(len(labels))}, "groups": {}, "notes": []}


# ── Pure renders ─────────────────────────────────────────────────────────────

def test_xlsx_letterhead_header_formats_subtotals_and_total():
    wb = openpyxl.load_workbook(io.BytesIO(export.to_xlsx(multi_result(), NOW)))
    assert wb.sheetnames == ["Rent Roll", "Summary"]
    ws = wb["Rent Roll"]
    cells = {ws.cell(r, c).value for r in range(1, 6) for c in range(1, 8)}
    assert "Rent Roll" in cells and "10/01/2026 through 10/20/2026" in cells
    assert "2 Facilities: GS Murrieta, GS Temecula" in cells
    assert "Generated on Tue 10/20/2026 at 3:05 PM PDT" in cells
    hdr = next(r for r in range(1, 10) if ws.cell(r, 1).value == "Facility")
    assert [ws.cell(hdr, c).value for c in range(1, 8)] == ["Facility", "Tenant", "Unit", "Move-In Date", "Monthly Rent", "Occupancy", "Units"]
    assert ws.cell(hdr, 1).font.bold and ws.cell(hdr, 1).fill.fgColor.rgb.endswith(export.CHARCOAL)
    assert ws.freeze_panes == f"A{hdr + 1}" and ws.print_title_rows == f"${hdr}:${hdr}"
    assert ws.auto_filter.ref.startswith(f"A{hdr}:G")
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1
    assert "&P of &N" in ws.oddFooter.right.text
    col_a = [ws.cell(r, 1).value for r in range(hdr + 1, ws.max_row + 1)]
    # Facilities in name order, each with a bar and a subtotal, then the grand total.
    assert col_a[:6] == ["GS Murrieta", "GS Murrieta", "GS Murrieta", "Total GS Murrieta", "GS Temecula", "GS Temecula"]
    sub = hdr + 1 + col_a.index("Total GS Murrieta")
    assert ws.cell(sub, 5).value == 75.0 and ws.cell(sub, 7).value == 2 and ws.cell(sub, 1).font.bold
    grand = hdr + 1 + col_a.index("Grand Total")
    assert ws.cell(grand, 5).value == 225.0 and ws.cell(grand, 5).number_format == export.MONEY_FMT
    assert ws.cell(grand, 5).border.bottom.style == "double"
    row = hdr + 2          # the first Murrieta line
    assert ws.cell(row, 5).value == 100.0 and ws.cell(row, 5).number_format == export.MONEY_FMT
    assert ws.cell(row, 4).value == datetime(2026, 1, 10) and ws.cell(row, 4).number_format == "mm/dd/yyyy"
    assert ws.cell(row, 6).value == 0.5 and ws.cell(row, 6).number_format == "0.0%"
    tricky = next(r for r in range(hdr, ws.max_row + 1) if str(ws.cell(r, 2).value).startswith("=cmd"))
    assert ws.cell(tricky, 2).data_type == "s"              # stored as text, never run as a formula
    assert export.PRIVACY in col_a and "A note about the rows." in col_a
    summary = wb["Summary"]
    assert summary.cell(7, 1).value == "Rent Roll" and summary.cell(7, 2).value == 225.0
    assert summary.cell(8, 2).value == pytest.approx(0.714) and summary.cell(8, 2).number_format == "0.0%"


def test_xlsx_one_facility_has_no_bars_and_takes_the_management_total_label():
    r = multi_result()
    r["rows"] = [x for x in r["rows"] if x["facility"] == "GS Temecula"]
    r["facilities"] = ["GS Temecula"]
    r["totals"] = {"facility": "All Facilities", "rent": 15000, "count": 2}
    ws = openpyxl.load_workbook(io.BytesIO(export.to_xlsx(r, NOW))).active
    col_a = [ws.cell(i, 1).value for i in range(1, ws.max_row + 1)]
    assert "Total GS Temecula" not in col_a and "All Facilities" in col_a
    assert any(ws.cell(i, 3).value == "Facility: GS Temecula" for i in range(1, 6))


def test_xlsx_detail_sheet_for_drill_down_groups():
    r = multi_result()
    r["rows"][0]["_sub"] = "s1"
    r["groups"] = {"s1": {"columns": [col("date", "Date", "date"), col("amount", "Amount", "money", True)],
                          "rows": [{"date": "2026-10-03", "amount": 1500}, {"date": "2026-10-04", "amount": 500}]}}
    wb = openpyxl.load_workbook(io.BytesIO(export.to_xlsx(r, NOW)))
    ws = wb["Detail"]
    values = [[ws.cell(i, c).value for c in (1, 2)] for i in range(1, ws.max_row + 1)]
    assert ["Date", "Amount"] in values and ["Total", 20.0] in values
    assert [datetime(2026, 10, 3), 15.0] in values


def test_pdf_renders_with_pages_and_handles_twenty_columns():
    pdf = export.to_pdf(multi_result(), NOW)
    assert pdf.startswith(b"%PDF") and pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") + pdf.count(b"/Type/Page") >= 1
    wide = export.to_pdf(wide_result(300), NOW)
    assert wide.startswith(b"%PDF")
    try:
        import fitz          # PyMuPDF, when the machine has it: read the text back
    except ImportError:
        return
    doc = fitz.open(stream=wide, filetype="pdf")
    assert doc.page_count > 1
    first, last = (" ".join(doc[i].get_text().split()) for i in (0, doc.page_count - 1))
    assert "Accounts Receivable" in first and "Reversals" in first and f"Page 1 of {doc.page_count}" in first
    assert "Accounts Receivable" in last and "Reversals" in last           # letterhead and column header on every page
    assert "Total" in last and export.PRIVACY in last
    for page in doc:                                                     # nothing runs off the page
        for x0, y0, x1, y1, *_ in page.get_text("words"):
            assert x1 <= page.rect.width - 20
    text = fitz.open(stream=export.to_pdf(multi_result(), NOW), filetype="pdf")[0].get_text()
    assert "GS Murrieta" in text and "Total GS Murrieta" in text and "Grand Total" in text
    assert "$100.00" in text and "($25.00)" in text and "RENT ROLL" in text and "$225.00" in text


def test_pdf_caps_huge_reports_and_says_so(monkeypatch):
    monkeypatch.setattr(export, "MAX_PDF_ROWS", 50)
    pdf = export.to_pdf(wide_result(80), NOW)
    try:
        import fitz
    except ImportError:
        return
    doc = fitz.open(stream=pdf, filetype="pdf")
    assert "the XLSX download has every line" in "".join(p.get_text() for p in doc).replace("\n", " ")


def test_csv_is_flat_with_dollars_us_dates_and_a_total():
    text = export.to_csv(multi_result())
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == ["Facility", "Tenant", "Unit", "Move-In Date", "Monthly Rent", "Occupancy", "Units"]
    assert rows[1] == ["GS Temecula", "Ada Sample", "A1", "03/01/2025", "150.00", "85.7", "1"]
    assert rows[2][1].startswith("'=")                                   # formula text defused
    assert rows[-1] == ["Total", "", "", "", "225.00", "", "4"]


def test_builder_results_map_and_render():
    built = {"columns": [{"key": "facility", "label": "Facility", "type": "text"}, {"key": "rent", "label": "Sum of Rent", "type": "money"},
                         {"key": "per", "label": "Rent per Sq Ft", "type": "num"}, {"key": "month", "label": "Month", "type": "text"}],
             "rows": [{"facility": "GS Temecula", "rent": 35000, "per": 1.75, "month": "2026-10"}],
             "totals": {"rent": 35000, "per": 1.75}, "dataset": "Units", "period": None, "facility": "All Facilities", "asOf": "2026-10-15"}
    r = export.from_builder(built, "Rent per Sq Ft")
    assert r["title"] == "Rent per Sq Ft" and export.period_line(r) == "As of 10/15/2026"
    ws = openpyxl.load_workbook(io.BytesIO(export.to_xlsx(r, NOW))).active
    flat = [ws.cell(i, c).value for i in range(1, ws.max_row + 1) for c in range(1, 5)]
    assert 350.0 in flat and "10/2026" in flat
    assert export.to_pdf(r, NOW).startswith(b"%PDF")


def test_filename_and_period_lines():
    r = multi_result()
    assert export.filename(r, "xlsx") == "Rent Roll - 2 Facilities - 10-01-2026 to 10-20-2026.xlsx"
    r["period"] = None
    r["facilities"] = None
    r["facility"] = "GS Temecula"
    assert export.filename(r, "pdf") == "Rent Roll - GS Temecula - as of 10-20-2026.pdf"
    assert export.period_line({"period": {"from": "2026-10-05", "to": "2026-10-05"}}) == "10/05/2026"
    assert export.text_of("2026-10-05T10:00:00Z", "text") == "10/05/2026"


# ── Endpoints (the seeded rows, Cubby faked) ─────────────────────────────────

@pytest.fixture
def client(monkeypatch):
    import cubby
    _seed()
    monkeypatch.setenv("CUBBY_API_KEY", "test-key-not-real")
    monkeypatch.setattr(cubby, "_facility_today", lambda fac: TODAY)
    monkeypatch.setattr(ops_reports, "_cubby", FakeCubby())
    ops_reports._ledger_cache.clear()
    return _client()


def _no_secrets(text: str):
    for s in SECRETS:
        assert s not in text, f"{s} leaked"


def test_export_endpoint_xlsx_for_two_facilities(client):
    r = client.get("/ops-reports/export/rent-roll", params={"format": "xlsx", "facility_ids": "fac_a,fac_b"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'attachment; filename="Rent Roll - ' in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    text = " ".join(str(ws.cell(i, c).value) for i in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1))
    assert "Ada Sample" in text and "Cy Park" in text and "Total" in text
    _no_secrets(text)


def test_export_endpoint_pdf_csv_one_facility_and_errors(client):
    pdf = client.get("/ops-reports/export/accounts-receivable", params={"format": "pdf", "facility_id": "fac_a"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == "application/pdf"
    c = client.get("/ops-reports/export/payment-detail", params={"format": "csv", "facility_ids": "fac_a",
                                                                  "start": "2026-10-01", "end": "2026-10-15"})
    assert c.status_code == 200 and c.headers["content-type"].startswith("text/csv")
    body = c.content.decode("utf-8-sig")
    assert "10/03/2026" in body and "150.00" in body
    _no_secrets(body)
    assert client.get("/ops-reports/export/rent-roll", params={"format": "docx"}).status_code == 422
    assert client.get("/ops-reports/export/nope").status_code == 404
    assert client.get("/ops-reports/export/inventory-activity").status_code == 422
    assert client.get("/ops-reports/export/rent-roll", params={"options": "not json"}).status_code == 422


def test_builder_export_endpoint(client):
    r = client.post("/ops-reports/builder/export", params={"format": "xlsx"}, json={"spec": _spec(), "name": "Rent per Sq Ft"})
    assert r.status_code == 200, r.text
    assert "Rent per Sq Ft" in r.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "Rent per Sq Ft"
    bad = client.post("/ops-reports/builder/export", json={"spec": {"dataset": "units", "measures": [{"id": "x", "agg": "sum", "field": "ssn"}]}})
    assert bad.status_code == 422


def test_run_report_shim_cuts_rows_when_run_takes_one_facility(client):
    result = export.run_report("rent-roll", ["fac_b"], "", "")
    assert {r["_f"] for r in result["rows"]} == {"fac_b"}
    both = export.run_report("rent-roll", ["fac_a", "fac_b"], "", "")
    assert {r["_f"] for r in both["rows"]} == {"fac_a", "fac_b"}
    assert date.fromisoformat(both["asOf"]) == TODAY
