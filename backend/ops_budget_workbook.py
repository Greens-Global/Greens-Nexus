"""Operations > Ops Budget: the workbook loader (Oct 21 2026).

One Ops Budget workbook per facility (Ops_Budget_Report_Template.xlsx):

  Setup        B3 facility name, B4 first month of the budget year, B5 first
               month of history
  SS Actual    header row 5, one row per month from row 6 (col A = month),
  VS Actual    one metric per column
  SS Budget    the same columns, 12 rows for the budget year
  VS Budget
  SS/VS Report, SS/VS Trend, Instructions - formulas and text, NOT read (the
  server recalculates everything)

Columns are read BY HEADER TEXT, never by position, matched to the metric
labels (and a few aliases) ignoring case, spacing and punctuation. A header we
do not know, and a metric with no column, are reported back to the uploader -
never guessed. A blank cell is NULL (a dash on screen); a typed 0 is zero.
Percentages are never typed: a column whose header is a calculated rate is
ignored and said so. A re-upload upserts every month in the file.

Upload is a two-step: preview (parse, nothing saved) then commit (the same
parse, then upsert + a load-log row that keeps the file). Admins only - the
router checks. Blocking (openpyxl, DB): plain `def` endpoints run it in
FastAPI's threadpool.
"""
from __future__ import annotations

import hashlib
import io
import re
import uuid
from datetime import date, datetime, timedelta

import ops_budget as ob

MAX_BYTES = 10 * 1024 * 1024
HEADER_SCAN_ROWS = 15
HEADER_ROW = 5               # where the template puts it; found by its "Month" cell either way


class WorkbookError(Exception):
    """The file cannot be loaded at all. The message is shown to the uploader."""


# ── Matching text ─────────────────────────────────────────────────────────────

_UNIT_TOKENS = [r"\$\s*/\s*mo(nth)?\b", r"/\s*mo(nth)?\b", r"\bper month\b", r"\(\s*\$\s*\)", r"\(\s*#\s*\)",
                r"\(\s*count\s*\)", r"\(\s*(f|s)\s*\)", r"\$", r"#", r"\bin dollars\b", r"\bdollars\b"]


def norm(text) -> str:
    """'Avg standard rate $/mo' -> 'avgstandardrate'; 'Past due 90+ days' -> 'pastdue90+days'."""
    s = str(text or "").lower().replace("&", " and ")
    for pat in _UNIT_TOKENS:
        s = re.sub(pat, " ", s)
    s = s.replace("square feet", "sq ft").replace("square foot", "sq ft").replace("average ", "avg ")
    return re.sub(r"[^a-z0-9+]", "", s)


def _header_index() -> tuple[dict, dict]:
    """({normalized header: input metric key}, {normalized header: rate key})."""
    inputs, rates = {}, {}
    for m in ob.METRICS:
        target = rates if m["kind"] == "rate" else inputs
        for text in [m["label"], *m["aliases"]]:
            target.setdefault(norm(text), m["key"])
    return inputs, rates


INPUT_HEADERS, RATE_HEADERS = _header_index()

SHEETS = {
    ("self_storage", "actual"): {"ssactual", "selfstorageactual", "selfstorageactuals", "ssactuals"},
    ("vehicle", "actual"): {"vsactual", "vehiclestorageactual", "vehicleactual", "vsactuals", "vehiclestorageactuals"},
    ("self_storage", "budget"): {"ssbudget", "selfstoragebudget"},
    ("vehicle", "budget"): {"vsbudget", "vehiclestoragebudget", "vehiclebudget"},
}
SHEET_LABEL = {("self_storage", "actual"): "SS Actual", ("vehicle", "actual"): "VS Actual",
               ("self_storage", "budget"): "SS Budget", ("vehicle", "budget"): "VS Budget"}


# ── Cells ─────────────────────────────────────────────────────────────────────

_MONTH_FORMATS = ("%b %Y", "%B %Y", "%Y-%m", "%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%b-%y", "%b %y", "%b-%Y", "%B-%Y",
                  "%m-%Y", "%Y/%m", "%m/%d/%y", "%d-%b-%y", "%d-%b-%Y")


def month_of(cell) -> str | None:
    """A month cell -> 'YYYY-MM', or None when it is not a month."""
    if cell is None:
        return None
    if isinstance(cell, datetime):
        return ob.ym(cell.date())
    if isinstance(cell, date):
        return ob.ym(cell)
    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
        if 20000 < cell < 80000:                       # an Excel serial date left unformatted
            return ob.ym(date(1899, 12, 30) + timedelta(days=int(cell)))
        return None
    text = str(cell).strip()
    if not text:
        return None
    text = re.sub(r"\s+", " ", text.replace(".", "").replace("'", " ").replace("’", " ")).strip()
    for fmt in _MONTH_FORMATS:
        try:
            return ob.ym(datetime.strptime(text, fmt).date())
        except ValueError:
            continue
    return None


def number_of(cell):
    """A value cell -> float, None for a blank. Raises ValueError for text."""
    if cell is None:
        return None
    if isinstance(cell, bool):
        raise ValueError("true/false")
    if isinstance(cell, (int, float)):
        return float(cell)
    if isinstance(cell, (datetime, date)):
        raise ValueError("a date")
    text = str(cell).strip()
    if text in ("", "-", "--", "—", "–", "n/a", "N/A", "na", "NA"):
        return None
    neg = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace("$", "").replace(",", "").replace(" ", "")
    if text.endswith("%"):
        raise ValueError("a percentage")
    v = float(text)
    return -v if neg else v


# ── Reading ───────────────────────────────────────────────────────────────────

def _open(raw: bytes):
    if not raw:
        raise WorkbookError("The file is empty.")
    if len(raw) > MAX_BYTES:
        raise WorkbookError("The file is larger than 10 MB. Upload the Ops Budget workbook only.")
    try:
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - openpyxl raises many kinds for a bad file
        raise WorkbookError("This is not an Excel workbook (.xlsx) we can read.") from e


def _rows(ws, limit: int = 400) -> list[tuple]:
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= limit:
            break
        out.append(tuple(row))
    return out


def _cell(rows, r: int, c: int):
    """1-based row and column, like Excel."""
    if r - 1 < len(rows) and c - 1 < len(rows[r - 1]):
        return rows[r - 1][c - 1]
    return None


def read_setup(rows) -> dict:
    """Setup B3 / B4 / B5. When B3 is empty, a 'Facility' label in column A
    points at its value instead."""
    name, start, history = _cell(rows, 3, 2), _cell(rows, 4, 2), _cell(rows, 5, 2)
    if not name:
        for r in range(1, min(len(rows), 12) + 1):
            label = norm(_cell(rows, r, 1))
            if label in ("facility", "facilityname", "site", "sitename", "property"):
                name = _cell(rows, r, 2)
            elif label.startswith("firstmonthofbudget") or label in ("budgetyearstart", "fiscalyearstart"):
                start = start or _cell(rows, r, 2)
            elif label.startswith("firstmonthofhistory") or label == "historystart":
                history = history or _cell(rows, r, 2)
    return {"name": str(name or "").strip(), "budgetYearStart": month_of(start) or "", "historyStart": month_of(history) or ""}


def read_metric_sheet(rows, label: str, segment: str) -> dict:
    """One Actual / Budget tab -> {"months": {YYYY-MM: {metric: value|None}},
    "found", "unknown", "missing", "ignored", "warnings"}."""
    head_at = None
    if norm(_cell(rows, HEADER_ROW, 1)) == "month":
        head_at = HEADER_ROW
    else:
        for r in range(1, min(len(rows), HEADER_SCAN_ROWS) + 1):
            if norm(_cell(rows, r, 1)) == "month":
                head_at = r
                break
    if head_at is None:
        raise WorkbookError(f"The {label} tab has no header row: we look for a row whose first cell says Month (row {HEADER_ROW} in the template).")
    header = rows[head_at - 1]
    cols: dict[int, str] = {}
    unknown, ignored, warnings = [], [], []
    seen: dict[str, int] = {}
    for c, text in enumerate(header[1:], start=1):
        if text is None or not str(text).strip():
            continue
        n = norm(text)
        key = INPUT_HEADERS.get(n)
        if key is None:
            if n in RATE_HEADERS:
                ignored.append(str(text).strip())
            else:
                unknown.append(str(text).strip())
            continue
        if ob.BY_KEY[key]["vehicleOnly"] and segment != "vehicle":
            ignored.append(str(text).strip())
            continue
        if key in seen:
            warnings.append(f"{label}: two columns read as {ob.BY_KEY[key]['label']} - both were skipped.")
            cols.pop(seen[key], None)
            continue
        seen[key] = c
        cols[c] = key
    dupes = {w.split(" read as ")[1].split(" - ")[0] for w in warnings if " read as " in w}
    found = [k for k in cols.values()]
    expected = [m["key"] for m in ob.metrics_for(segment) if m["kind"] != "rate"]
    missing = [ob.BY_KEY[k]["label"] for k in expected if k not in found and ob.BY_KEY[k]["label"] not in dupes]

    months: dict[str, dict] = {}
    for i, row in enumerate(rows[head_at:], start=head_at + 1):
        if not row or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        mo = month_of(row[0] if row else None)
        if mo is None:
            if any(v is not None and str(v).strip() for v in row[1:]):
                warnings.append(f"{label} row {i}: column A is not a month ({str(row[0])[:30]!r}), so the row was skipped.")
            continue
        if mo in months:
            warnings.append(f"{label} row {i}: {ob.month_label(mo)} appears twice - the later row was used.")
        vals = {}
        for c, key in cols.items():
            cell = row[c] if c < len(row) else None
            try:
                vals[key] = number_of(cell)
            except ValueError as e:
                warnings.append(f"{label} row {i}, {ob.BY_KEY[key]['label']}: {str(cell)[:30]!r} is {e}, not a number - left blank.")
                vals[key] = None
        months[mo] = vals
    return {"months": months, "found": found, "unknown": unknown, "missing": missing, "ignored": ignored,
            "warnings": warnings}


def _sheet_key(name: str):
    n = norm(name)
    for key, names in SHEETS.items():
        if n in names:
            return key
    return None


def parse(raw: bytes) -> dict:
    """The whole workbook -> what would be loaded. Raises WorkbookError."""
    wb = _open(raw)
    try:
        setup = None
        sheets = {}
        for ws in wb.worksheets:
            n = norm(ws.title)
            if n == "setup" and setup is None:
                setup = read_setup(_rows(ws, 30))
                continue
            key = _sheet_key(ws.title)
            if key and key not in sheets:
                sheets[key] = read_metric_sheet(_rows(ws), SHEET_LABEL[key], key[0])
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    if setup is None:
        raise WorkbookError("The workbook has no Setup tab (facility name in B3, first month of the budget year in B4).")
    if not sheets:
        raise WorkbookError("The workbook has none of the SS Actual, VS Actual, SS Budget or VS Budget tabs.")
    return {"setup": setup, "sheets": sheets}


def match_facility(name: str, facilities: list[dict]) -> dict | None:
    n = norm(name)
    if not n:
        return None
    hits = [f for f in facilities if norm(f["shortName"]) and norm(f["shortName"]) in n]
    return max(hits, key=lambda f: len(f["shortName"])) if len(hits) >= 1 else None


def build_preview(db, raw: bytes, file_name: str, facility_id: str = "") -> dict:
    """Parse and check against the facility list. Nothing is saved."""
    got = parse(raw)
    facs = ob.facilities(db)
    setup = got["setup"]
    errors, warnings = [], []
    fac = match_facility(setup["name"], facs)
    chosen = next((f for f in facs if f["id"] == facility_id), None) if facility_id else None
    if facility_id and chosen is None:
        errors.append("That facility is not one of the Ops Budget sites.")
    if chosen and fac and chosen["id"] != fac["id"]:
        errors.append(f"The Setup tab says {setup['name']}, which is not {chosen['shortName']}. Check you picked the right file.")
    fac = fac or chosen
    if fac is None:
        errors.append(f"We could not tell which site this workbook is for. Setup B3 says {setup['name']!r}; "
                      f"it should name one of {', '.join(f['shortName'] for f in facs)}." if setup["name"] else
                      "Setup B3 is empty - it should name the facility.")
    bys = setup["budgetYearStart"]
    if any(k[1] == "budget" for k in got["sheets"]) and not bys:
        errors.append("Setup B4 (first month of the budget year) is not a month, so the budget cannot be placed.")

    sheets_out, values = [], []
    budget_end = ob.add_months(bys, 11) if bys else ""
    for key in sorted(got["sheets"], key=lambda k: list(SHEETS).index(k)):
        seg, scenario = key
        s = got["sheets"][key]
        kept = {}
        for mo, vals in sorted(s["months"].items()):
            if scenario == "budget" and bys and not (bys <= mo <= budget_end):
                s["warnings"].append(f"{SHEET_LABEL[key]}: {ob.month_label(mo)} is outside the budget year "
                                     f"({ob.month_label(bys)} - {ob.month_label(budget_end)}) and was skipped.")
                continue
            kept[mo] = vals
        filled = [mo for mo, vals in kept.items() if any(v is not None for v in vals.values())]
        for mo, vals in kept.items():
            for k, v in vals.items():
                values.append((fac["id"] if fac else "", seg, mo, scenario, k, v))
        sheets_out.append({
            "sheet": SHEET_LABEL[key], "segment": seg, "scenario": scenario,
            "months": len(kept), "monthsWithData": len(filled),
            "first": min(filled) if filled else "", "last": max(filled) if filled else "",
            "values": sum(1 for vals in kept.values() for v in vals.values() if v is not None),
            "columns": [ob.BY_KEY[k]["label"] for k in s["found"]],
            "unknown": s["unknown"], "missing": s["missing"], "ignored": s["ignored"], "warnings": s["warnings"],
        })
        warnings += s["warnings"]
    for key in SHEETS:
        if key not in got["sheets"]:
            warnings.append(f"There is no {SHEET_LABEL[key]} tab, so nothing was read for it.")
    return {
        "fileName": file_name, "facility": fac, "setup": setup, "sheets": sheets_out,
        "errors": errors, "warnings": warnings, "values": values,
        "unknownHeaders": sorted({h for s in sheets_out for h in s["unknown"]}),
        "missingHeaders": sorted({f"{s['sheet']}: {h}" for s in sheets_out for h in s["missing"]}),
    }


def public(preview: dict) -> dict:
    """The preview without the raw value list."""
    return {k: v for k, v in preview.items() if k != "values"}


def save(db, preview: dict, raw: bytes, by: str) -> dict:
    """Upsert the preview's values and log the load (the file kept). Locks
    the facility row so two uploads of one site queue instead of racing."""
    from ops_budget_models import OpsBudgetFacility, OpsBudgetLoad
    if preview["errors"]:
        raise WorkbookError(preview["errors"][0])
    fid = preview["facility"]["id"]
    fac = db.query(OpsBudgetFacility).filter(OpsBudgetFacility.id == fid).with_for_update().first()
    if fac is None:
        raise WorkbookError("That facility is not one of the Ops Budget sites.")
    stamp = ob.now_iso()
    load_id = str(uuid.uuid4())
    setup = preview["setup"]
    if setup["budgetYearStart"]:
        fac.budget_year_start = setup["budgetYearStart"]
        fac.fiscal_start_month = int(setup["budgetYearStart"][5:7])
    if setup["historyStart"]:
        fac.history_start = setup["historyStart"]
    fac.updated_at, fac.updated_by = stamp, by
    written = ob.upsert(db, preview["values"], source="workbook", by=by, load_id=load_id, stamp=stamp)
    months = sorted({v[2] for v in preview["values"]})
    db.add(OpsBudgetLoad(
        id=load_id, facility_id=fid, kind="workbook", file_name=(preview.get("fileName") or "")[:200],
        file_sha256=hashlib.sha256(raw).hexdigest(), file_size=len(raw), file_bytes=raw, months=months,
        rows_written=written, unknown_headers=preview["unknownHeaders"], missing_headers=preview["missingHeaders"],
        warnings=preview["warnings"][:200], status="ok", error="", loaded_by=by, loaded_at=stamp))
    db.commit()
    return {"loadId": load_id, "rowsWritten": written, "months": len(months)}
