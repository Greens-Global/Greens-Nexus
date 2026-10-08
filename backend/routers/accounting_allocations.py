"""Accounting > Allocations (Neil, 10/01: "allocations journal entry from
people and done monthly").

Every month payroll is paid by one entity per company while the people
worked for several: the allocation entry moves each person's wages to the
entities they worked for, in proportion to the hours they clocked at each
work site. The basis is Time Clock's hours by location (the same
`byLocation` roll-up the By-location report and the IIF export read, off
`_compute_timecard` / `_fixed_card`); the cost is that month's wages from
the same payroll card - never a second pay computation. For each person:
debit each entity's wage expense account by its share, credit the paying
entity's allocation clearing account for the whole; a salaried person with
no site hours is left on the paying entity (flagged, not allocated).

Which work site is which Intacct entity and account, and which company pays
through which entity and clearing account, is the MAPPING - saved under
nexus_settings key `accounting_allocations_map` and edited on the screen; a
site or company without a mapping shows as unmapped so nothing is exported
half-coded. A run (accounting_allocation_runs) keeps the whole preview as
it was, so the export (Intacct GL import CSV, shared layout in
intacct_gl.py, or Excel) can be produced again unchanged. Nothing is posted
anywhere.

Wages for everyone are payroll figures, so this takes the Full level on the
Accounting grant AND a fresh step-up, like the payroll card itself.
"""
import io
import json
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

import intacct_gl
import models
from auth import get_current_user, require_module_grant
from database import get_db
from routers.accounting import _display_name
from routers.stepup import require_stepup

router = APIRouter(
    prefix="/accounting/allocations",
    tags=["Accounting"],
    dependencies=[Depends(require_module_grant("accounting", "full"))],
)

MAP_KEY = "accounting_allocations_map"
_DEFAULT_MAP = {"journal": "GJ", "sites": {}, "companies": {}, "offsite": {"entity": "", "account": ""}}
_MAX_MAP_CHARS = 60000
NO_SITE = "__none__"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _month_range(month: str) -> tuple[str, str]:
    """"2026-09" -> ("2026-09-01", "2026-09-30")."""
    try:
        first = datetime.strptime((month or "")[:7], "%Y-%m")
    except ValueError:
        raise HTTPException(status_code=400, detail="month must be YYYY-MM")
    nxt = first.replace(year=first.year + 1, month=1) if first.month == 12 else first.replace(month=first.month + 1)
    last = nxt - timedelta(days=1)
    return first.strftime("%Y-%m-%d"), last.strftime("%Y-%m-%d")


def _mmyyyy(month: str) -> str:
    return f"{month[5:7]}/{month[0:4]}"


# ── The mapping ──────────────────────────────────────────────────────────────

def _clean_map(raw: Any) -> dict:
    out = {"journal": "GJ", "sites": {}, "companies": {}, "offsite": {"entity": "", "account": ""}}
    if not isinstance(raw, dict):
        return out
    out["journal"] = (str(raw.get("journal") or "GJ").strip() or "GJ")[:20]
    for key in ("sites", "companies"):
        src = raw.get(key) or {}
        if isinstance(src, dict):
            for k, v in src.items():
                if isinstance(v, dict) and k:
                    out[key][str(k)[:80]] = {f: str(v.get(f) or "").strip()[:40] for f in ("entity", "account", "wageAccount")}
    off = raw.get("offsite") or {}
    if isinstance(off, dict):
        out["offsite"] = {f: str(off.get(f) or "").strip()[:40] for f in ("entity", "account")}
    return out


def _read_map(db: Session) -> dict:
    row = db.query(models.NexusSetting).filter(models.NexusSetting.key == MAP_KEY).first()
    if not row or not row.value:
        return dict(_DEFAULT_MAP, sites={}, companies={}, offsite=dict(_DEFAULT_MAP["offsite"]))
    try:
        return _clean_map(json.loads(row.value))
    except ValueError:
        return dict(_DEFAULT_MAP, sites={}, companies={}, offsite=dict(_DEFAULT_MAP["offsite"]))


class MapBody(BaseModel):
    journal: str = "GJ"
    sites: dict[str, dict] = {}
    companies: dict[str, dict] = {}
    offsite: dict = {}


@router.get("/map")
def get_map(user: dict = Depends(get_current_user), db: Session = Depends(get_db)):
    """The mapping plus every work site and company it can name."""
    sites = [{"id": s.id, "name": s.name} for s in db.query(models.HrWorkSite).order_by(models.HrWorkSite.name).all()]
    companies = [{"id": c.id, "name": c.name} for c in db.query(models.HrEntity).order_by(models.HrEntity.name).all()]
    return {"map": _read_map(db), "sites": sites, "companies": companies}


@router.put("/map")
def put_map(body: MapBody, user: dict = Depends(get_current_user), db: Session = Depends(get_db)):
    clean = _clean_map(body.model_dump())
    text = json.dumps(clean)
    if len(text) > _MAX_MAP_CHARS:
        raise HTTPException(status_code=400, detail="That mapping is too large.")
    row = db.query(models.NexusSetting).filter(models.NexusSetting.key == MAP_KEY).first()
    if not row:
        row = models.NexusSetting(key=MAP_KEY)
        db.add(row)
    row.value, row.updated_by, row.updated_at = text, user["email"], _now()
    db.commit()
    return {"map": clean}


# ── The preview ──────────────────────────────────────────────────────────────

def _split(total: float, weights: list[int]) -> list[float]:
    """`total` across `weights` to the cent, the last share taking the
    rounding so the parts add up to the whole."""
    whole = sum(weights)
    if whole <= 0 or not weights:
        return [0.0] * len(weights)
    out = [round(total * w / whole, 2) for w in weights[:-1]]
    out.append(round(total - sum(out), 2))
    return out


def build_preview(db: Session, month: str, mapping: dict, cards: dict[str, dict], people: list[dict]) -> dict:
    """The entry from already-computed payroll cards. `people` rows carry
    email, name, payType, company (HrEntity id), companyName, department,
    employeeId; `cards[email]` is the payroll card (totals.totalPay,
    byLocation). Pure, so the math is testable without Time Clock."""
    start, end = _month_range(month)
    sites_map, companies_map, offsite = mapping.get("sites") or {}, mapping.get("companies") or {}, mapping.get("offsite") or {}
    label = _mmyyyy(month)
    out_people, unmapped_sites, unmapped_companies = [], {}, {}
    by_entry: dict[str, dict] = {}

    def entry_for(cid: str, cname: str, paying: dict) -> dict:
        key = cid or "__nocompany__"
        if key not in by_entry:
            by_entry[key] = {"journal": mapping.get("journal") or "GJ", "date": end, "reference_no": "",
                             "description": f"Payroll allocation {label} - {cname or 'No company'}"[:intacct_gl.DESC_MAX],
                             "companyId": cid, "companyName": cname, "payingEntity": paying.get("entity") or "", "lines": []}
        return by_entry[key]

    for p in people:
        card = cards.get(p["email"]) or {}
        totals = card.get("totals") or {}
        wages = round(float(totals.get("totalPay") or 0), 2)
        locs = [x for x in (card.get("byLocation") or []) if (x.get("workedMin") or 0) > 0]
        total_min = sum(int(x.get("workedMin") or 0) for x in locs)
        if wages == 0 and total_min == 0:
            continue
        cid = p.get("company") or ""
        paying = companies_map.get(cid) or {}
        if not paying.get("entity") or not paying.get("account"):
            unmapped_companies[cid or NO_SITE] = p.get("companyName") or "No company"
        amounts = _split(wages, [int(x.get("workedMin") or 0) for x in locs]) if total_min else []
        shares = []
        entry = entry_for(cid, p.get("companyName") or "", paying)
        if total_min:
            for x, amt in zip(locs, amounts):
                sid = x.get("workSiteId") or ""
                m = sites_map.get(sid) if sid else offsite
                m = m or {}
                mapped = bool(m.get("entity") and m.get("account"))
                if not mapped:
                    unmapped_sites[sid or NO_SITE] = x.get("workSite") or "No location"
                pct = round(100.0 * int(x.get("workedMin") or 0) / total_min, 1)
                shares.append({"workSiteId": sid, "workSite": x.get("workSite") or "No location", "workedMin": int(x.get("workedMin") or 0),
                               "share": pct, "amount": amt, "entity": m.get("entity") or "", "account": m.get("account") or "", "mapped": mapped})
                entry["lines"].append({"type": "debit", "email": p["email"], "name": p["name"], "employee_id": p.get("employeeId") or "",
                                       "acct_no": m.get("account") or "", "location_id": m.get("entity") or "", "dept_id": p.get("department") or "",
                                       "source_entity": paying.get("entity") or "", "debit": amt, "credit": None, "mapped": mapped,
                                       "memo": f"{p['name']} - {x.get('workSite') or 'No location'} - {int(x.get('workedMin') or 0) / 60:.2f} h of {total_min / 60:.2f} h ({pct}%)"})
        else:
            # Salaried with no site hours this month: stays with the payer.
            mapped = bool(paying.get("entity") and paying.get("wageAccount"))
            shares.append({"workSiteId": "", "workSite": "Not allocated (no site hours)", "workedMin": 0, "share": 100.0, "amount": wages,
                           "entity": paying.get("entity") or "", "account": paying.get("wageAccount") or "", "mapped": mapped})
            entry["lines"].append({"type": "debit", "email": p["email"], "name": p["name"], "employee_id": p.get("employeeId") or "",
                                   "acct_no": paying.get("wageAccount") or "", "location_id": paying.get("entity") or "", "dept_id": p.get("department") or "",
                                   "source_entity": paying.get("entity") or "", "debit": wages, "credit": None, "mapped": mapped,
                                   "memo": f"{p['name']} - not allocated (no site hours in {label})"})
        entry["lines"].append({"type": "credit", "email": p["email"], "name": p["name"], "employee_id": p.get("employeeId") or "",
                               "acct_no": paying.get("account") or "", "location_id": paying.get("entity") or "", "dept_id": p.get("department") or "",
                               "source_entity": paying.get("entity") or "", "debit": None, "credit": wages,
                               "mapped": bool(paying.get("entity") and paying.get("account")),
                               "memo": f"{p['name']} - payroll allocation {label}"})
        out_people.append({"email": p["email"], "name": p["name"], "payType": p.get("payType") or "hourly", "currency": card.get("currency") or p.get("currency") or "USD",
                           "company": cid, "companyName": p.get("companyName") or "", "department": p.get("department") or "",
                           "wages": wages, "workedMin": total_min, "sites": shares})

    entries = []
    for e in by_entry.values():
        if not e["lines"]:
            continue
        for i, line in enumerate(e["lines"]):
            line["line_no"] = i + 1
        entries.append(e)
    debits = round(sum(float(line.get("debit") or 0) for e in entries for line in e["lines"]), 2)
    credits = round(sum(float(line.get("credit") or 0) for e in entries for line in e["lines"]), 2)
    unmapped_lines = sum(1 for e in entries for line in e["lines"] if not line.get("mapped"))
    currencies = sorted({p["currency"] for p in out_people})
    return {"month": month, "start": start, "end": end, "people": out_people, "entries": entries,
            "totals": {"wages": round(sum(p["wages"] for p in out_people), 2), "debits": debits, "credits": credits, "people": len(out_people), "unmappedLines": unmapped_lines},
            "unmapped": {"sites": [{"id": k, "name": v} for k, v in unmapped_sites.items()], "companies": [{"id": k, "name": v} for k, v in unmapped_companies.items()]},
            "currencies": currencies, "mapping": mapping}


def _people_and_cards(db: Session, month: str) -> tuple[list[dict], dict[str, dict]]:
    """Everyone on the payroll roster for the month with their card - the
    SAME engine as the timecard screen (Time Clock owns it)."""
    from routers.timeclock import _compute_timecard, _fixed_card, _pay_type, _team_rows
    start, end = _month_range(month)
    rows = _team_rows(db, start, end, only_emails=None, include_fixed=True)
    emps = {(e.work_email or "").lower(): e for e in db.query(models.NexusEmployee).all() if e.work_email}
    companies = {c.id: c.name for c in db.query(models.HrEntity).all()}
    rates = {r.employee_email: r for r in db.query(models.PayrollRate).all()}
    people, cards = [], {}
    for r in rows:
        em = r["email"]
        emp = emps.get(em)
        cid = (emp.company if emp else "") or ""
        fixed = _pay_type(db, em) == "fixed"
        card = _fixed_card(db, em, start) if fixed else _compute_timecard(db, em, start, end)
        rate = rates.get(em)
        cards[em] = card
        people.append({"email": em, "name": r["name"], "payType": "fixed" if fixed else "hourly", "company": cid,
                       "companyName": companies.get(cid, ""), "department": (emp.department if emp else "") or "",
                       "employeeId": (emp.employee_code if emp else "") or "", "currency": (getattr(rate, "currency", None) or "USD") if rate else "USD"})
    return people, cards


@router.get("/preview")
def preview(month: str, user: dict = Depends(get_current_user), _su: dict = Depends(require_stepup), db: Session = Depends(get_db)):
    """The allocation entry for a month as it stands: basis per person, the
    lines, what is unmapped. Nothing is saved."""
    mapping = _read_map(db)
    people, cards = _people_and_cards(db, month)
    return build_preview(db, month, mapping, cards, people)


# ── Runs ─────────────────────────────────────────────────────────────────────

class RunBody(BaseModel):
    month: str
    entity: str = ""
    preview: dict


def _run_out(r: models.AccountingAllocationRun, db: Session, full: bool = False) -> dict:
    data = r.lines or {}
    out = {"id": r.id, "month": r.month, "entity": r.entity, "by": r.run_by, "byName": _display_name(r.run_by, db), "at": r.run_at,
           "totals": data.get("totals") or {}, "people": len(data.get("people") or []), "entries": len(data.get("entries") or [])}
    if full:
        out["preview"] = data
    return out


@router.get("/runs")
def list_runs(user: dict = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.query(models.AccountingAllocationRun).order_by(models.AccountingAllocationRun.run_at.desc()).limit(200).all()
    return {"runs": [_run_out(r, db) for r in rows]}


@router.post("/runs", status_code=201)
def save_run(body: RunBody, user: dict = Depends(get_current_user), _su: dict = Depends(require_stepup), db: Session = Depends(get_db)):
    """Keep a preview as a run. The entries must balance; unmapped lines are
    refused so no half-coded file reaches Intacct."""
    _month_range(body.month)
    pv = body.preview or {}
    entries = pv.get("entries") or []
    if not entries:
        raise HTTPException(status_code=400, detail="There is nothing to keep - the preview has no lines.")
    if any(not line.get("mapped") for e in entries for line in (e.get("lines") or [])):
        raise HTTPException(status_code=400, detail="Map every work site and company first - some lines have no entity or account.")
    if not intacct_gl.balanced(entries):
        raise HTTPException(status_code=400, detail="The entry does not balance.")
    if len(json.dumps(pv)) > 4_000_000:
        raise HTTPException(status_code=400, detail="That run is too large to keep.")
    row = models.AccountingAllocationRun(id=str(uuid.uuid4()), month=body.month[:7], entity=(body.entity or "").strip()[:40],
                                         run_by=user["email"], run_at=_now(), lines=pv)
    db.add(row)
    db.commit()
    return _run_out(row, db, full=True)


def _run(db: Session, run_id: str) -> models.AccountingAllocationRun:
    row = db.query(models.AccountingAllocationRun).filter(models.AccountingAllocationRun.id == run_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="That run is gone.")
    return row


@router.get("/runs/{run_id}")
def get_run(run_id: str, user: dict = Depends(get_current_user), _su: dict = Depends(require_stepup), db: Session = Depends(get_db)):
    return _run_out(_run(db, run_id), db, full=True)


@router.get("/runs/{run_id}/export.csv")
def export_run_csv(run_id: str, user: dict = Depends(get_current_user), _su: dict = Depends(require_stepup), db: Session = Depends(get_db)):
    """The run as an Intacct General Ledger import file (the shared layout)."""
    row = _run(db, run_id)
    body = intacct_gl.to_csv((row.lines or {}).get("entries") or [])
    fname = f"Intacct GL Import - Payroll Allocation - {row.month}.csv"
    return StreamingResponse(iter([body]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.get("/runs/{run_id}/export.xlsx")
def export_run_xlsx(run_id: str, user: dict = Depends(get_current_user), _su: dict = Depends(require_stepup), db: Session = Depends(get_db)):
    """The run as a workbook: the import sheet as Intacct reads it, and the
    basis behind it (each person's hours by site and the share of wages)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    row = _run(db, run_id)
    data = row.lines or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Intacct Import"
    for n, r in enumerate(intacct_gl.to_rows(data.get("entries") or [])):
        ws.append(r if n == 0 else [float(c) if i in (12, 13) and c else c for i, c in enumerate(r)])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Basis")
    ws2.append(["Person", "Pay Type", "Company", "Department", "Wages", "Currency", "Work Site", "Hours", "Share %", "Amount", "Entity", "Account"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for p in data.get("people") or []:
        for s in p.get("sites") or []:
            ws2.append([p.get("name"), p.get("payType"), p.get("companyName"), p.get("department"), float(p.get("wages") or 0), p.get("currency"),
                        s.get("workSite"), round(float(s.get("workedMin") or 0) / 60, 2), float(s.get("share") or 0), float(s.get("amount") or 0), s.get("entity"), s.get("account")])
    ws2.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    fname = f"Payroll Allocation - {row.month}.xlsx"
    return StreamingResponse(iter([buf.getvalue()]), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@router.delete("/runs/{run_id}", status_code=204)
def delete_run(run_id: str, user: dict = Depends(get_current_user), db: Session = Depends(get_db)):
    """A run kept by mistake can be removed; the Intacct side is untouched
    because nothing was ever posted."""
    row = _run(db, run_id)
    db.delete(row)
    db.commit()
    return None

