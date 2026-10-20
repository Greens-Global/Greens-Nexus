"""Operations > Reports: the downloadable files (XLSX, PDF, CSV) - Oct 20 2026.

Cubby's Reports page is nothing but downloads; ours is interactive, and the
files it hands out have to look better than Cubby's too (Visesh: "even in
xlsx look at the UI formatting"). One report result (ops_reports.run, or a
Report Builder result mapped by `from_builder`) in, bytes out:

  to_xlsx   letterhead (logo, title, period, facilities, generated-on), a
            branded header row (wrapped labels, frozen, filtered, repeated on
            every printed page), real Excel dates (mm/dd/yyyy), accounting
            money with "-" for zero, percents, a facility bar and subtotal per
            facility when several are in one report, a grand total, Notes, a
            "Summary" sheet of headline figures and a "Detail" sheet with the
            drill-down lines. Landscape, fit to width, "Page X of Y".
  to_pdf    landscape letter; the letterhead on EVERY page, the column header
            repeated on each page, summary tiles on page 1, zebra rows, money
            right-aligned with "$" and "-" for zero, facility bars, subtotals,
            grand total, notes, "Page X of Y". Many columns (the AR
            roll-forward has ~20) shrink the type and wrap the labels; long
            text is cut with an ellipsis instead of spilling over.
  to_csv    flat data with a Facility column and a Total line, dollars and
            MM/DD/YYYY; text a spreadsheet would run as a formula is quoted.

Optional result keys (another branch adds them to ops_reports.run): `facilities`
(short names) and `summary` ([{label, value, type}]); everything works without.

Privacy: the rows are already whitelisted by ops_reports (tenant name, unit
and money only). Nothing here adds a field - only the columns the result
declares are written, and link keys (_f, _c, _u, _sub) never are.
"""
from __future__ import annotations

import csv
import io
import os
import re
from collections import OrderedDict
from datetime import date, datetime
from zoneinfo import ZoneInfo

TZ = "America/Los_Angeles"          # Greens Storage facilities are in California
PRIVACY = "Tenant name, unit and money only: no phone, email or address."
MAX_PDF_ROWS = 5000                 # the PDF is the print-ready lighter copy; the XLSX has every line
MAX_DETAIL_ROWS = 20000

_LOGO_PATHS = (
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "greens-storage-logo.png"),
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                 "frontend", "public", "assets", "branding", "greens-global-logo.png"),
)

# Brand: the Greens Storage charcoal and lime.
CHARCOAL = "54565B"
LIME = "8DC63F"
INK = "1F2328"
DIM = "6B7280"
BAR = "EEF5E1"           # facility group bar
SUB = "F2F3F5"           # subtotal
GRAND = "E4E6E9"         # grand total
ZEBRA = "F8F9FA"

MONEY_FMT = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
INT_FMT = '#,##0_);(#,##0);"-"_)'
NUM_FMT = '#,##0.00_);(#,##0.00);"-"_)'
PCT_FMT = '0.0%'
DATE_FMT = 'mm/dd/yyyy'

NUMERIC = {"money", "int", "num", "area", "pct"}
_ISO_DAY = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_ISO_MONTH = re.compile(r"^\d{4}-\d{2}$")
_HIDDEN = {"_f", "_c", "_u", "_sub"}


# ── Shared helpers ────────────────────────────────────────────────────────────

def logo_path() -> str:
    for p in _LOGO_PATHS:
        if os.path.exists(p):
            return p
    return ""


def us_date(iso) -> str:
    s = str(iso or "")
    return f"{s[5:7]}/{s[8:10]}/{s[:4]}" if len(s) >= 10 and s[4] == "-" else ""


def _as_date(v):
    s = str(v or "")[:10]
    if not _ISO_DAY.match(s):
        return None
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None


def _now() -> datetime:
    return datetime.now(ZoneInfo(TZ))


def _clock(dt: datetime) -> str:
    return dt.strftime("%I:%M %p").lstrip("0")


def generated_line(now: datetime | None = None) -> str:
    """'Generated on Tue 10/20/2026 at 3:05 PM PDT' - US date, 12-hour clock."""
    now = now or _now()
    return f"Generated on {now:%a} {now:%m/%d/%Y} at {_clock(now)} {now.strftime('%Z')}".strip()


def stamp_line(iso: str) -> str:
    """A UTC iso stamp as Pacific 'MM/DD/YYYY h:mm AM PDT'."""
    try:
        dt = datetime.fromisoformat(str(iso).replace("Z", "+00:00")).astimezone(ZoneInfo(TZ))
    except (ValueError, TypeError):
        return ""
    return f"{dt:%m/%d/%Y} {_clock(dt)} {dt.strftime('%Z')}"


def period_line(result: dict) -> str:
    p = result.get("period")
    if p and p.get("from") and p.get("to"):
        span = (f"{us_date(p['from'])} through {us_date(p['to'])}" if p["from"] != p["to"] else us_date(p["from"]))
        label = p.get("label")
        return f"{label}: {span}" if label and label != "Custom" else span
    return f"As of {us_date(result.get('asOf'))}" if result.get("asOf") else ""


def facility_names(result: dict) -> list[str]:
    names = result.get("facilities")
    if isinstance(names, list) and names:
        return [str(n) for n in names]
    one = str(result.get("facility") or "").strip()
    return [one] if one else []


def facilities_line(result: dict) -> str:
    names = facility_names(result)
    if not names:
        return ""
    if len(names) == 1:
        return names[0] if names[0] == "All Facilities" else f"Facility: {names[0]}"
    return f"{len(names)} Facilities: " + ", ".join(names)


def filename(result: dict, ext: str) -> str:
    title = result.get("title") or "Ops Report"
    names = facility_names(result)
    fac = names[0] if len(names) == 1 else f"{len(names)} Facilities" if names else ""
    p = result.get("period")
    if p and p.get("from") and p.get("to"):
        when = f"{us_date(p['from'])} to {us_date(p['to'])}".replace("/", "-")
    else:
        when = f"as of {us_date(result.get('asOf'))}".replace("/", "-") if result.get("asOf") else ""
    stem = " - ".join(x for x in (title, fac, when) if x)
    stem = re.sub(r"[^A-Za-z0-9 ._()&-]", "", stem.replace("/", "-"))
    stem = re.sub(r"\s+", " ", stem).strip()[:150] or "Ops Report"
    return f"{stem}.{ext}"


def columns(result: dict) -> list[dict]:
    return [c for c in (result.get("columns") or []) if c.get("key") not in _HIDDEN]


def _first_text(cols) -> str:
    for c in cols:
        if c.get("type") not in NUMERIC:
            return c["key"]
    return cols[0]["key"] if cols else ""


def sections(result: dict):
    """[(facility or None, rows)]. Rows are split by facility when the report
    has a Facility column, several facilities, and more than one line for at
    least one of them (one line per facility is a summary already)."""
    rows = result.get("rows") or []
    keys = {c["key"] for c in columns(result)}
    if "facility" not in keys:
        return [(None, rows)]
    by = OrderedDict()
    for r in rows:
        by.setdefault(str(r.get("facility") or ""), []).append(r)
    if len(by) < 2 or len(rows) <= len(by):
        return [(None, rows)]
    return [(k, by[k]) for k in sorted(by, key=lambda s: s.lower())]


def subtotal(cols, rows) -> dict:
    out = {}
    for c in cols:
        if c.get("total") and c.get("type") in NUMERIC and c.get("type") != "pct":
            out[c["key"]] = sum((r.get(c["key"]) or 0) for r in rows if isinstance(r.get(c["key"]), (int, float)))
    return out


def has_numbers(values: dict, cols) -> bool:
    return any(isinstance(values.get(c["key"]), (int, float)) and not isinstance(values.get(c["key"]), bool)
               for c in cols if c.get("type") in NUMERIC)


def text_of(v, t: str) -> str:
    """A value as the PDF / email prints it."""
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if t == "money":
        if not isinstance(v, (int, float)):
            return str(v)
        if round(v) == 0:
            return "-"
        s = f"${abs(v) / 100:,.2f}"
        return f"({s})" if v < 0 else s
    if t == "pct":
        return f"{v:.1f}%" if isinstance(v, (int, float)) else str(v)
    if t in ("int", "area"):
        return f"{round(v):,}" if isinstance(v, (int, float)) else str(v)
    if t == "num":
        if isinstance(v, (int, float)):
            return f"{v:,.2f}" if not float(v).is_integer() else f"{v:,.0f}"
        return str(v)
    s = str(v)
    if _ISO_DAY.match(s[:10]) and (len(s) == 10 or s[10] == "T"):
        return us_date(s)
    if _ISO_MONTH.match(s):
        return f"{s[5:7]}/{s[:4]}"
    return s


def notes_of(result: dict) -> list[str]:
    out = [str(n) for n in (result.get("notes") or []) if str(n).strip()]
    stamp = stamp_line(result.get("fetchedAt") or result.get("syncedAt") or "")
    if stamp:
        out.append(f"Figures from Cubby as of {stamp}.")
    out.append(PRIVACY)
    return out


# ── Report Builder results ────────────────────────────────────────────────────

def from_builder(result: dict, name: str = "") -> dict:
    """ops_report_builder.run() output in the shape the exporters take."""
    cols = [{"key": c["key"], "label": c["label"], "type": c["type"], "total": c["type"] in NUMERIC and c["type"] != "pct"}
            for c in result.get("columns") or []]
    period = result.get("period")
    about = f"Report Builder: {result.get('dataset') or ''}".strip()
    return {"key": "builder", "title": (name or "").strip() or result.get("dataset") or "Report Builder",
            "category": "Report Builder", "about": about, "source": "nexus",
            "period": {"from": period["from"], "to": period["to"], "label": period.get("label")} if period else None,
            "asOf": result.get("asOf") or "", "facility": result.get("facility") or "All Facilities",
            "available": True, "columns": cols, "rows": result.get("rows") or [],
            "totals": result.get("totals") or {}, "groups": {},
            "notes": [about + "." if about else "", "Built from the rows Nexus syncs from Cubby every 15 minutes."]}


# ── CSV ───────────────────────────────────────────────────────────────────────

def _csv_cell(v, t: str) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if t == "money" and isinstance(v, (int, float)):
        return f"{v / 100:.2f}"
    if t in ("int", "area") and isinstance(v, (int, float)):
        return str(round(v))
    if t == "pct" and isinstance(v, (int, float)):
        return f"{v:.1f}"
    if t == "num" and isinstance(v, (int, float)):
        return f"{round(v, 4):g}"
    s = text_of(v, "text")
    if s and s[0] in "=+-@\t\r" and not re.match(r"^-?\d", s):
        s = "'" + s    # a name a spreadsheet would run as a formula
    return s


def to_csv(result: dict) -> str:
    cols = columns(result)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in cols])
    for r in result.get("rows") or []:
        w.writerow([_csv_cell(r.get(c["key"]), c.get("type", "text")) for c in cols])
    totals = result.get("totals") or {}
    if cols and has_numbers(totals, cols):
        first = _first_text(cols)
        line = []
        for c in cols:
            v = totals.get(c["key"])
            if c["key"] == first:
                line.append(str(v) if isinstance(v, str) and v else "Total")
            else:
                line.append(_csv_cell(v, c.get("type", "text")) if c.get("type") in NUMERIC else "")
        w.writerow(line)
    return buf.getvalue()


# ── XLSX ──────────────────────────────────────────────────────────────────────

def _xl_value(v, t: str):
    """(value, number_format or None)."""
    if v is None or v == "":
        return None, None
    if isinstance(v, bool):
        return ("Yes" if v else "No"), None
    if t == "money" and isinstance(v, (int, float)):
        return round(v / 100, 2), MONEY_FMT
    if t == "pct" and isinstance(v, (int, float)):
        return v / 100, PCT_FMT
    if t in ("int", "area") and isinstance(v, (int, float)):
        return v, INT_FMT
    if t == "num" and isinstance(v, (int, float)):
        return v, NUM_FMT
    if t == "date" or isinstance(v, str):
        d = _as_date(v) if (t == "date" or (isinstance(v, str) and _ISO_DAY.match(v))) else None
        if d:
            return d, DATE_FMT
    if isinstance(v, (int, float)):
        return v, None
    return text_of(v, "text"), None


def _display_len(v, t: str) -> int:
    if v is None or v == "":
        return 0
    if t == "money":
        return len(text_of(v, t)) + 3
    if t == "date":
        return 10
    return len(text_of(v, t))


def _widths(cols, rows) -> list[float]:
    out = []
    sample = rows[:3000]
    for c in cols:
        t = c.get("type", "text")
        longest_word = max((len(w) for w in str(c["label"]).split()), default=4)
        label_need = max(longest_word, (len(c["label"]) + 1) // 2)     # labels wrap onto two lines
        data_need = max((_display_len(r.get(c["key"]), t) for r in sample), default=0)
        w = max(label_need, data_need) + 2
        if t == "money":
            w = max(w, 14)
        lo, hi = (9, 22) if t in NUMERIC or t == "date" else (8, 42)
        out.append(float(min(max(w, lo), hi)))
    return out


def to_xlsx(result: dict, now: datetime | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    cols = columns(result)
    title = result.get("title") or "Ops Report"
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Report"
    ws.sheet_view.showGridLines = False
    ncol = max(len(cols), 1)
    last = get_column_letter(ncol)

    font = "Calibri"
    thin = Side(style="thin", color="D9DCE1")
    hdr_fill = PatternFill("solid", fgColor=CHARCOAL)
    bar_fill = PatternFill("solid", fgColor=BAR)
    sub_fill = PatternFill("solid", fgColor=SUB)
    grand_fill = PatternFill("solid", fgColor=GRAND)
    zebra = PatternFill("solid", fgColor=ZEBRA)

    widths = _widths(cols, result.get("rows") or [])
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Letterhead: the logo top-left over rows 1-4, the text block to its right.
    logo = logo_path()
    text_col = 1
    if logo:
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(logo)
            img.height = 62
            img.width = round(62 * 700 / 286)
            ws.add_image(img, "A1")
            need, acc, text_col = img.width / 7 + 1, 0.0, ncol + 1
            for i, w in enumerate(widths, start=1):
                if acc >= need:
                    text_col = i
                    break
                acc += w
            if text_col > ncol:      # a narrow report: the text goes under the logo
                text_col = 1
        except Exception:  # noqa: BLE001 - a missing image library degrades the letterhead, not the file
            logo = ""
    block = [(title, Font(name=font, size=15, bold=True, color=INK), 24),
             (period_line(result), Font(name=font, size=11, color=INK), 17),
             (facilities_line(result), Font(name=font, size=10, color=DIM), 15),
             (generated_line(now), Font(name=font, size=9, italic=True, color=DIM), 15)]
    row0 = 1
    if logo and text_col == 1:
        ws.row_dimensions[1].height = 50
        row0 = 2
    if not logo:
        ws.cell(row=1, column=1, value="GREENS STORAGE").font = Font(name=font, size=9, bold=True, color=LIME)
        row0 = 2
    for i, (text, f, h) in enumerate(block):
        cell = ws.cell(row=row0 + i, column=text_col, value=text)
        cell.font = f
        ws.row_dimensions[row0 + i].height = h
    hdr = row0 + len(block) + 1
    ws.row_dimensions[hdr - 1].height = 8

    # Header row.
    wrap_needed = any(len(c["label"]) > widths[i] - 2 for i, c in enumerate(cols))
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=hdr, column=i, value=c["label"])
        cell.font = Font(name=font, size=10, bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="right" if c.get("type") in NUMERIC else "left", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=LIME))
    ws.row_dimensions[hdr].height = 30 if wrap_needed else 20

    body = Font(name=font, size=10, color=INK)
    bold = Font(name=font, size=10, bold=True, color=INK)
    first = _first_text(cols)
    r = hdr + 1
    zebra_on = False

    def write_values(values: dict, f, fill=None, border=None, label=""):
        nonlocal r
        for i, c in enumerate(cols, start=1):
            t = c.get("type", "text")
            if c["key"] == first and label:
                v, fmt = label, None
            elif label and t not in NUMERIC:
                v, fmt = None, None
            else:
                v, fmt = _xl_value(values.get(c["key"]), t)
            cell = ws.cell(row=r, column=i, value=v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"          # text, never a formula
            if fmt:
                cell.number_format = fmt
            cell.font = f
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
            if t in NUMERIC:
                cell.alignment = Alignment(horizontal="right")
        r += 1

    secs = sections(result)
    grouped = secs[0][0] is not None
    for name, rows in secs:
        if grouped:
            for i in range(1, ncol + 1):
                ws.cell(row=r, column=i).fill = bar_fill
            cell = ws.cell(row=r, column=1, value=name or "No Facility")
            cell.font = Font(name=font, size=10.5, bold=True, color=INK)
            ws.row_dimensions[r].height = 18
            r += 1
            zebra_on = False
        for row in rows:
            write_values(row, body, zebra if zebra_on else None)
            zebra_on = not zebra_on
        if grouped and cols:
            write_values(subtotal(cols, rows), bold, sub_fill, Border(top=thin), label=f"Total {name}")
    last_data = r - 1

    totals = result.get("totals") or {}
    if cols and has_numbers(totals, cols):
        label = totals.get(first) if isinstance(totals.get(first), str) and totals.get(first) else (
            "Grand Total" if grouped else "Total")
        write_values(totals, Font(name=font, size=10.5, bold=True, color=INK), grand_fill,
                     Border(top=Side(style="thin", color=CHARCOAL), bottom=Side(style="double", color=CHARCOAL)), label=label)
        ws.row_dimensions[r - 1].height = 19

    # Notes.
    r += 1
    n_rows = len(result.get("rows") or [])
    ws.cell(row=r, column=1, value=f"{n_rows:,} Line{'s' if n_rows != 1 else ''}").font = Font(name=font, size=9, color=DIM)
    r += 2
    ws.cell(row=r, column=1, value="Notes").font = Font(name=font, size=10, bold=True, color=INK)
    r += 1
    for n in notes_of(result):
        cell = ws.cell(row=r, column=1, value=n)
        cell.font = Font(name=font, size=9, italic=(n == PRIVACY), color=DIM)
        if n.startswith("="):
            cell.data_type = "s"
        r += 1

    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    if cols and last_data > hdr:
        ws.auto_filter.ref = f"A{hdr}:{last}{last_data}"
    ws.print_title_rows = f"{hdr}:{hdr}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top, ws.page_margins.bottom = 0.6, 0.6
    ws.oddHeader.left.text = "&B" + title.replace("&", "&&")
    ws.oddHeader.right.text = (period_line(result) or "").replace("&", "&&")
    ws.oddFooter.left.text = ("Greens Storage - " + (facilities_line(result) or "")).replace("&", "&&")[:200]
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_options.horizontalCentered = True

    _summary_sheet(wb, result, now)
    _detail_sheet(wb, result)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_head(ws, result, sub: str, now, width_cols: int):
    from openpyxl.styles import Font
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=result.get("title") or "Ops Report").font = Font(name="Calibri", size=14, bold=True, color=INK)
    ws.cell(row=2, column=1, value=sub).font = Font(name="Calibri", size=11, color=INK)
    ws.cell(row=3, column=1, value=" | ".join(x for x in (period_line(result), facilities_line(result)) if x)).font = Font(
        name="Calibri", size=10, color=DIM)
    ws.cell(row=4, column=1, value=generated_line(now)).font = Font(name="Calibri", size=9, italic=True, color=DIM)


def _summary_sheet(wb, result: dict, now):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    summary = [s for s in (result.get("summary") or []) if isinstance(s, dict) and s.get("label")]
    if not summary:
        return
    ws = wb.create_sheet("Summary")
    _sheet_head(ws, result, "Summary", now, 2)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 20
    for i, label in enumerate(("Figure", "Value"), start=1):
        c = ws.cell(row=6, column=i, value=label)
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=CHARCOAL)
        c.border = Border(bottom=Side(style="medium", color=LIME))
        c.alignment = Alignment(horizontal="right" if i == 2 else "left")
    for n, s in enumerate(summary):
        row = 7 + n
        a = ws.cell(row=row, column=1, value=str(s["label"]))
        a.font = Font(name="Calibri", size=11, color=INK)
        v, fmt = _xl_value(s.get("value"), s.get("type") or "num")
        b = ws.cell(row=row, column=2, value=v)
        if isinstance(v, str) and v.startswith("="):
            b.data_type = "s"
        b.font = Font(name="Calibri", size=11, bold=True, color=INK)
        b.alignment = Alignment(horizontal="right")
        if fmt:
            b.number_format = fmt
        if n % 2:
            for c in (a, b):
                c.fill = PatternFill("solid", fgColor=ZEBRA)


def _detail_sheet(wb, result: dict):
    """The drill-down lines (Cubby's PDF has them inline; a sheet is easier to
    filter). Skipped when there are none or too many."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    groups = result.get("groups") or {}
    parents = [r for r in (result.get("rows") or []) if r.get("_sub") and r["_sub"] in groups]
    if not parents:
        return
    total = sum(len(groups[p["_sub"]].get("rows") or []) for p in parents)
    if total == 0 or total > MAX_DETAIL_ROWS:
        return
    ws = wb.create_sheet("Detail")
    _sheet_head(ws, result, "Detail: the lines behind each row", None, 0)
    pcols = [c for c in columns(result) if c.get("type") not in NUMERIC][:3]
    widths: dict[int, float] = {}
    r = 6
    for p in parents:
        g = groups[p["_sub"]]
        gcols = [c for c in (g.get("columns") or []) if c.get("key") not in _HIDDEN]
        label = " - ".join(text_of(p.get(c["key"]), c.get("type", "text")) for c in pcols if p.get(c["key"]))
        cell = ws.cell(row=r, column=1, value=label or "Detail")
        cell.font = Font(name="Calibri", size=10.5, bold=True, color=INK)
        for i in range(1, max(len(gcols), 1) + 1):
            ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=BAR)
        r += 1
        for i, c in enumerate(gcols, start=1):
            h = ws.cell(row=r, column=i, value=c["label"])
            h.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            h.fill = PatternFill("solid", fgColor=CHARCOAL)
            h.alignment = Alignment(horizontal="right" if c.get("type") in NUMERIC else "left")
        r += 1
        for w_i, w in enumerate(_widths(gcols, g.get("rows") or []), start=1):
            widths[w_i] = max(widths.get(w_i, 0), w)
        for row in g.get("rows") or []:
            for i, c in enumerate(gcols, start=1):
                v, fmt = _xl_value(row.get(c["key"]), c.get("type", "text"))
                cell = ws.cell(row=r, column=i, value=v)
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"
                if fmt:
                    cell.number_format = fmt
                cell.font = Font(name="Calibri", size=10, color=INK)
            r += 1
        sub = subtotal(gcols, g.get("rows") or [])
        if has_numbers(sub, gcols):
            for i, c in enumerate(gcols, start=1):
                v, fmt = _xl_value(sub.get(c["key"]), c.get("type", "text")) if c.get("type") in NUMERIC else (
                    ("Total" if i == 1 else None), None)
                cell = ws.cell(row=r, column=i, value=v)
                cell.font = Font(name="Calibri", size=10, bold=True, color=INK)
                cell.border = Border(top=Side(style="thin", color="D9DCE1"))
                if fmt:
                    cell.number_format = fmt
            r += 1
        r += 1
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.oddFooter.right.text = "Page &P of &N"


# ── PDF ───────────────────────────────────────────────────────────────────────

def _esc(s) -> str:
    return str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _fit(text: str, width: float, font: str, size: float) -> str:
    from reportlab.pdfbase.pdfmetrics import stringWidth
    if stringWidth(text, font, size) <= width:
        return text
    ell = "…"
    lo, hi = 0, len(text)
    while lo < hi:
        mid = (lo + hi + 1) // 2
        if stringWidth(text[:mid].rstrip() + ell, font, size) <= width:
            lo = mid
        else:
            hi = mid - 1
    return (text[:lo].rstrip() + ell) if lo else ""


def _draw_wordmark(cv, x, y_top, height):
    """The Greens Storage mark when the logo file is not on this server."""
    from reportlab.lib import colors
    cv.saveState()
    size = height * 0.55
    cv.setFont("Helvetica-Bold", size)
    cv.setFillColor(colors.HexColor("#" + CHARCOAL))
    cv.drawString(x, y_top - size, "G")
    gw = cv.stringWidth("G", "Helvetica-Bold", size)
    cv.setFillColor(colors.HexColor("#" + LIME))
    cv.drawString(x + gw, y_top - size, "REENS")
    cv.setFillColor(colors.HexColor("#" + CHARCOAL))
    cv.setFont("Helvetica", height * 0.26)
    cv.drawString(x + 1, y_top - height * 0.95, "S T O R A G E")
    cv.restoreState()


def to_pdf(result: dict, now: datetime | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas as _canvas
    from reportlab.platypus import (BaseDocTemplate, Frame, LongTable, PageTemplate, Paragraph, Spacer, Table,
                                    TableStyle)

    hx = lambda h: colors.HexColor("#" + h)  # noqa: E731
    page_w, page_h = landscape(letter)
    margin = 32
    head_h = 74
    foot_h = 30
    avail = page_w - 2 * margin
    title = result.get("title") or "Ops Report"
    gen = generated_line(now)
    per = period_line(result)
    facs = facilities_line(result)
    logo = logo_path()

    def furniture(cv, doc):
        cv.saveState()
        top = page_h - 26
        lh = 40
        if logo:
            try:
                cv.drawImage(logo, margin, top - lh, width=lh * 700 / 286, height=lh, mask="auto")
            except Exception:  # noqa: BLE001
                _draw_wordmark(cv, margin, top, lh)
        else:
            _draw_wordmark(cv, margin, top, lh)
        right = page_w - margin
        cv.setFillColor(hx(DIM))
        cv.setFont("Helvetica", 7.5)
        cv.drawRightString(right, top - 6, gen)
        cv.setFillColor(hx(INK))
        cv.setFont("Helvetica-Bold", 14)
        cv.drawRightString(right, top - 22, _fit(title, avail * 0.62, "Helvetica-Bold", 14))
        cv.setFont("Helvetica", 9)
        cv.drawRightString(right, top - 34, per)
        cv.setFillColor(hx(DIM))
        cv.setFont("Helvetica", 8)
        cv.drawRightString(right, top - 45, _fit(facs, avail * 0.62, "Helvetica", 8))
        cv.setStrokeColor(hx(LIME))
        cv.setLineWidth(1.6)
        cv.line(margin, page_h - head_h - 4, page_w - margin, page_h - head_h - 4)
        cv.setStrokeColor(hx("D9DCE1"))
        cv.setLineWidth(0.5)
        cv.line(margin, foot_h, page_w - margin, foot_h)
        cv.setFillColor(hx(DIM))
        cv.setFont("Helvetica", 7)
        cv.drawString(margin, foot_h - 11, "Greens Storage | Nexus Operations Reports | " + PRIVACY)
        cv.restoreState()

    class Numbered(_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._pages)
            for i, state in enumerate(self._pages, start=1):
                self.__dict__.update(state)
                self.setFont("Helvetica", 7.5)
                self.setFillColor(hx(DIM))
                self.drawRightString(page_w - margin, foot_h - 11, f"Page {i} of {n}")
                super().showPage()
            super().save()

    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=(page_w, page_h), leftMargin=margin, rightMargin=margin,
                          topMargin=head_h + 12, bottomMargin=foot_h + 8, title=title, author="Greens Storage - Nexus",
                          subject=per)
    frame = Frame(margin, foot_h + 8, avail, page_h - head_h - 12 - foot_h - 8, id="body",
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="page", frames=[frame], onPage=furniture)])

    story = []
    note_style = ParagraphStyle("note", fontName="Helvetica", fontSize=8, leading=11, textColor=hx(DIM))

    # Summary tiles (page 1).
    summary = [s for s in (result.get("summary") or []) if isinstance(s, dict) and s.get("label")][:12]
    if summary:
        per_row = min(6, len(summary))
        tile_w = avail / per_row
        lab = ParagraphStyle("tl", fontName="Helvetica", fontSize=7, leading=9, textColor=hx(DIM))
        val = ParagraphStyle("tv", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=hx(INK))
        cells = [[Paragraph(_esc(str(s["label"]).upper()), lab), Paragraph(_esc(text_of(s.get("value"), s.get("type") or "num")
                                                                               or "-"), val)] for s in summary]
        grid = []
        for i in range(0, len(cells), per_row):
            line = cells[i:i + per_row]
            line += [""] * (per_row - len(line))
            grid.append(line)
        tiles = Table(grid, colWidths=[tile_w] * per_row)
        style = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 8),
                 ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]
        for ri, line in enumerate(grid):
            for ci, c in enumerate(line):
                if c:
                    style += [("BACKGROUND", (ci, ri), (ci, ri), hx("F6F8F3")),
                              ("LINEBEFORE", (ci, ri), (ci, ri), 2.5, hx(LIME)),
                              ("LINEBELOW", (ci, ri), (ci, ri), 2, colors.white)]
        tiles.setStyle(TableStyle(style))
        story += [tiles, Spacer(1, 10)]

    cols = columns(result)
    secs = sections(result)
    grouped = secs[0][0] is not None
    distinct = {str(r.get("facility") or "") for r in result.get("rows") or []}
    if "facility" in {c["key"] for c in cols} and (grouped or len(distinct) <= 1) and len(cols) > 1:
        cols = [c for c in cols if c["key"] != "facility"]      # the bar or the letterhead says it
    if not result.get("available", True) or not cols:
        for n in notes_of(result):
            story.append(Paragraph(_esc(n), note_style))
        doc.build(story or [Spacer(1, 1)], canvasmaker=Numbered)
        return buf.getvalue()

    # Rows, capped for print.
    flat: list[tuple[str, object, dict]] = []      # (kind, label, values)
    shown = 0
    capped = False
    for name, rows in secs:
        if capped:
            break
        if grouped:
            flat.append(("bar", name or "No Facility", {}))
        part = []
        for row in rows:
            if shown >= MAX_PDF_ROWS:
                capped = True
                break
            flat.append(("row", "", row))
            part.append(row)
            shown += 1
        if grouped and not capped:
            flat.append(("sub", f"Total {name}", subtotal(cols, rows)))
    totals = result.get("totals") or {}
    first = _first_text(cols)
    if has_numbers(totals, cols):
        lbl = totals.get(first) if isinstance(totals.get(first), str) and totals.get(first) else ("Grand Total" if grouped else "Total")
        flat.append(("grand", lbl, totals))

    # Type size and column widths: shrink until it fits the page.
    n = len(cols)
    size = 8.0 if n <= 8 else 7.5 if n <= 11 else 7.0 if n <= 14 else 6.5 if n <= 17 else 6.0
    pad = 3

    def natural(sz):
        out = []
        sample = [v for k, _, v in flat if k == "row"][:1500] + [v for k, _, v in flat if k in ("sub", "grand")]
        for c in cols:
            t = c.get("type", "text")
            data = max((stringWidth(text_of(r.get(c["key"]), t), "Helvetica", sz) for r in sample), default=0)
            word = max((stringWidth(w, "Helvetica-Bold", sz) for w in str(c["label"]).split()), default=10)
            want = max(data, word, stringWidth(str(c["label"]), "Helvetica-Bold", sz) / 2.6) + 2 * pad + 2
            if t not in NUMERIC and t != "date":
                want = min(want, 210)
            out.append(want)
        return out

    w = natural(size)
    while sum(w) > avail and size > 5.5:
        size -= 0.5
        w = natural(size)
    if sum(w) > avail:
        text_idx = [i for i, c in enumerate(cols) if c.get("type") not in NUMERIC and c.get("type") != "date"]
        over = sum(w) - avail
        shrinkable = sum(max(0, w[i] - 40) for i in text_idx)
        if shrinkable > 0:
            k = min(1.0, over / shrinkable)
            for i in text_idx:
                w[i] -= max(0, w[i] - 40) * k
        if sum(w) > avail:
            f = avail / sum(w)
            w = [x * f for x in w]
    else:
        extra = avail - sum(w)
        grow = [i for i, c in enumerate(cols) if c.get("type") not in NUMERIC] or list(range(n))
        base = sum(w[i] for i in grow)
        for i in grow:
            w[i] += extra * (w[i] / base)

    head_style = ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=size, leading=size + 1.5, textColor=colors.white)
    head_r = ParagraphStyle("thr", parent=head_style, alignment=2)
    data = [[Paragraph(_esc(c["label"]), head_r if c.get("type") in NUMERIC else head_style) for c in cols]]
    style = [
        ("FONT", (0, 0), (-1, -1), "Helvetica", size),
        ("TEXTCOLOR", (0, 1), (-1, -1), hx(INK)),
        ("BACKGROUND", (0, 0), (-1, 0), hx(CHARCOAL)),
        ("LINEBELOW", (0, 0), (-1, 0), 1.5, hx(LIME)),
        ("VALIGN", (0, 0), (-1, 0), "BOTTOM"),
        ("VALIGN", (0, 1), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), pad), ("RIGHTPADDING", (0, 0), (-1, -1), pad),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6),
        ("TOPPADDING", (0, 0), (-1, 0), 4), ("BOTTOMPADDING", (0, 0), (-1, 0), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, hx(ZEBRA)]),
    ]
    for i, c in enumerate(cols):
        if c.get("type") in NUMERIC:
            style.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    first_i = next((i for i, c in enumerate(cols) if c["key"] == first), 0)
    lead = next((i for i, c in enumerate(cols) if c.get("type") in NUMERIC), n)
    for kind, label, values in flat:
        ri = len(data)
        if kind == "bar":
            data.append([_fit(str(label), avail - 2 * pad, "Helvetica-Bold", size + 0.5)] + [""] * (n - 1))
            style += [("SPAN", (0, ri), (-1, ri)), ("BACKGROUND", (0, ri), (-1, ri), hx(BAR)),
                      ("FONT", (0, ri), (-1, ri), "Helvetica-Bold", size + 0.5),
                      ("TOPPADDING", (0, ri), (-1, ri), 4), ("BOTTOMPADDING", (0, ri), (-1, ri), 4)]
            continue
        bold = kind in ("sub", "grand")
        fname = "Helvetica-Bold" if bold else "Helvetica"
        line = []
        for i, c in enumerate(cols):
            t = c.get("type", "text")
            if bold and c["key"] == first:
                s = _fit(str(label), sum(w[first_i:lead]) - 2 * pad, fname, size + 0.5) if lead > first_i else str(label)
                line.append(s)
                continue
            elif bold and t not in NUMERIC:
                s = ""
            else:
                s = text_of(values.get(c["key"]), t)
            line.append(_fit(s, w[i] - 2 * pad, fname, size))
        data.append(line)
        if bold and lead - first_i > 1:
            style.append(("SPAN", (first_i, ri), (lead - 1, ri)))     # the label spans the leading text columns
        if kind == "sub":
            style += [("FONT", (0, ri), (-1, ri), "Helvetica-Bold", size), ("BACKGROUND", (0, ri), (-1, ri), hx(SUB)),
                      ("LINEABOVE", (0, ri), (-1, ri), 0.6, hx("C4C8CE"))]
        elif kind == "grand":
            style += [("FONT", (0, ri), (-1, ri), "Helvetica-Bold", size + 0.5), ("BACKGROUND", (0, ri), (-1, ri), hx(GRAND)),
                      ("LINEABOVE", (0, ri), (-1, ri), 1, hx(CHARCOAL)), ("LINEBELOW", (0, ri), (-1, ri), 1, hx(CHARCOAL)),
                      ("TOPPADDING", (0, ri), (-1, ri), 4), ("BOTTOMPADDING", (0, ri), (-1, ri), 4)]
    if len(data) == 1:
        data.append(["No lines for this period and facility."] + [""] * (n - 1))
        style += [("SPAN", (0, 1), (-1, 1)), ("TEXTCOLOR", (0, 1), (-1, 1), hx(DIM))]
    table = LongTable(data, colWidths=w, repeatRows=1)
    table.setStyle(TableStyle(style))
    story.append(table)

    story.append(Spacer(1, 10))
    total_rows = len(result.get("rows") or [])
    count = f"{total_rows:,} Line{'s' if total_rows != 1 else ''}"
    if capped:
        count += f". This print copy shows the first {MAX_PDF_ROWS:,}; the XLSX download has every line."
    story.append(Paragraph(_esc(count), note_style))
    story.append(Spacer(1, 6))
    story.append(Paragraph("<b>Notes</b>", ParagraphStyle("nh", parent=note_style, textColor=hx(INK), fontSize=8.5)))
    for note in notes_of(result):
        story.append(Paragraph(_esc(note), note_style))
    doc.build(story, canvasmaker=Numbered)
    return buf.getvalue()


# ── One door for the router and the schedules ─────────────────────────────────

MEDIA = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf",
         "csv": "text/csv; charset=utf-8"}


def render(result: dict, fmt: str) -> tuple[bytes, str, str]:
    """(bytes, media type, filename) for xlsx | pdf | csv."""
    fmt = (fmt or "xlsx").lower()
    if fmt == "pdf":
        data = to_pdf(result)
    elif fmt == "csv":
        data = ("﻿" + to_csv(result)).encode("utf-8")     # the BOM makes Excel read UTF-8
    elif fmt == "xlsx":
        data = to_xlsx(result)
    else:
        raise ValueError("A format is xlsx, pdf or csv")
    return data, MEDIA[fmt], filename(result, fmt)


def run_report(kind: str, facility_ids: list[str] | None = None, start: str = "", end: str = "",
               options: dict | None = None) -> dict:
    """ops_reports.run for one or several facilities, with report options.
    Blocking. Works whether or not ops_reports.run takes `facility_ids` /
    `options` yet (they land on another branch): the signature is read at
    call time. Without `facility_ids`, several facilities run as all of them
    and the rows are cut down to the ones asked for."""
    import inspect

    import ops_reports
    ids = [i for i in dict.fromkeys(str(x).strip() for x in (facility_ids or [])) if i]
    params = inspect.signature(ops_reports.run).parameters
    kw = {}
    if options:
        if "options" in params:
            kw["options"] = dict(options)
        else:
            fixed = {"kind", "facility_id", "facility_ids", "start", "end"}
            kw.update({k: v for k, v in options.items() if k in params and k not in fixed})
    if len(ids) <= 1:
        return ops_reports.run(kind, ids[0] if ids else "", start, end, **kw)
    if "facility_ids" in params:
        return ops_reports.run(kind, "", start, end, facility_ids=ids, **kw)
    result = ops_reports.run(kind, "", start, end, **kw)
    wanted = set(ids)
    rows = [r for r in result.get("rows") or [] if r.get("_f") in wanted]
    names = sorted({str(r.get("facility")) for r in rows if r.get("facility")})
    return {**result, "rows": rows, "totals": ops_reports._totals(result.get("columns") or [], rows),
            "facility": ", ".join(names) or "Selected Facilities", "facilities": names or None}
