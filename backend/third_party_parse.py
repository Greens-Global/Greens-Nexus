"""Read a third-party operator's report export (CubeSmart first) into rows
(Operations, Oct 20 2026). Pure: bytes in, dicts out - no database, no network.

Three reports, each read by its LABELS, never by column position, so a
re-ordered export still lands in the right fields:

  sre        CubeSmart's Qlik "SRE Summary" (qsreports.cubesmart.com): the
             Rental Activity and Revenue & AR tables, one column per month
             ("Oct-2026", "Sep-2026", ...), one row per metric. A long layout
             (one row per month, one column per metric) reads too.
  flash      SmartView's "Daily Flash" (smartviewweb.cubesmart.com/flash): the
             Net Rentals, Occupancy by Status, Revenue, Allowances and Net
             Revenue tables stacked on a sheet (as the screen shows them), or a
             long layout with one row per day.
  rent_roll  the Rent Roll sheet of CubeSmart's monthly Financials workbook (the
             one site_models/build/cubesmart_units.py reads), or a CSV of it.
             ONLY Store, Cube, Cube Dimensions and Cube Attribute are read -
             customer names, rates and dates in the other columns are never
             touched, so they never reach the database or an error message.

Excel (.xlsx) or CSV. Every report shows numbers the way Qlik exports them:
"2,926", "85.1%", "$1,215.13", "(1,227)", "-", "NA".
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Tuple

MAX_BYTES = 15 * 1024 * 1024
KINDS = ("sre", "flash", "rent_roll")

MONTHS = {m: i + 1 for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"))}


class ParseError(ValueError):
    """A file Nexus cannot read. The message is safe to show and to log: it
    never quotes a cell (a rent roll cell can be a customer's name)."""


# ── Cells ────────────────────────────────────────────────────────────────────

def norm(v: Any) -> str:
    """Label key: lower case, '%' kept as the word 'pct', everything else that
    is not a letter or digit folded to one space. "Tot % Rented - SqF" ->
    "tot pct rented sqf"; "MoveIns" -> "moveins"."""
    s = str(v if v is not None else "").strip().lower().replace("%", " pct ").replace("#", " number ")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def num(v: Any, pct: bool = False) -> Optional[float]:
    """A report cell as a number, or None for an empty / "-" / "NA" cell.
    `pct`: the value is a percentage on a 0-100 scale; a fraction (0.851, how
    Excel stores a cell formatted as a percent) is scaled up."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return round(x * 100, 4) if pct and -1.5 <= x <= 1.5 and x != 0 else x
    s = str(v).strip()
    if not s or s.upper() in ("-", "NA", "N/A", "--", "NULL", "#N/A"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    had_pct = "%" in s
    s = re.sub(r"[,$%()\s]", "", s)
    if not re.fullmatch(r"[-+]?\d+(\.\d+)?", s):
        return None
    x = float(s)
    if neg:
        x = -x
    if pct and not had_pct and -1.5 <= x <= 1.5 and x != 0:
        x = x * 100
    return x


def month_of(v: Any) -> Optional[str]:
    """A month header as YYYY-MM: "Oct-2026", "Oct 2026", "October 2026",
    "Oct-26", "2026-10", "10/2026", or an Excel date. None otherwise."""
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v or "").strip().lower()
    if not s:
        return None
    m = re.fullmatch(r"([a-z]{3})[a-z]*[\s\-/']+(\d{2}|\d{4})", s)
    if m and m.group(1) in MONTHS:
        y = int(m.group(2))
        y = y + 2000 if y < 100 else y
        return f"{y:04d}-{MONTHS[m.group(1)]:02d}"
    m = re.fullmatch(r"(\d{4})-(\d{1,2})", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    m = re.fullmatch(r"(\d{1,2})/(\d{4})", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return f"{int(m.group(2)):04d}-{int(m.group(1)):02d}"
    return None


def day_of(v: Any) -> Optional[str]:
    """A date cell as YYYY-MM-DD: an Excel date, "10/19/2026", "2026-10-19"."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(1)), int(m.group(2))).isoformat()
        except ValueError:
            return None
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return None
    return None


def as_of_in_text(v: Any) -> Optional[str]:
    """The day in a report title: "Daily Flash Landing Page as of 10/19/2026 10:28:40 AM"."""
    s = str(v or "")
    m = re.search(r"as\s+of\s+(\d{1,2}/\d{1,2}/\d{4}|\d{4}-\d{2}-\d{2})", s, re.I)
    return day_of(m.group(1)) if m else None


def store_of(v: Any) -> str:
    """A store cell as its number: 6992, "6992", "6992 - Austin", "Store 6992"."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    m = re.search(r"\d{3,6}", str(v or ""))
    return m.group(0) if m else ""


# ── Files -> grids ───────────────────────────────────────────────────────────

def read_sheets(raw: bytes, file_name: str = "") -> List[Tuple[str, List[List[Any]]]]:
    """[(sheet name, rows)] from an .xlsx or a .csv. Values only (formulas as
    their last computed value)."""
    if not raw:
        raise ParseError("The file is empty.")
    if len(raw) > MAX_BYTES:
        raise ParseError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    name = (file_name or "").lower()
    if raw[:2] == b"PK" or name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 - any unreadable workbook is the same answer
            raise ParseError("Could not open that workbook. Export it again as .xlsx or .csv.") from e
        out = []
        for ws in wb.worksheets:
            out.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
        wb.close()
        return out
    if name.endswith(".xls"):
        raise ParseError("Old .xls workbooks are not supported. Save it as .xlsx or .csv and upload again.")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [("CSV", [list(r) for r in csv.reader(io.StringIO(text), dialect)])]


def _cells(row: List[Any]) -> List[Tuple[int, Any]]:
    return [(i, c) for i, c in enumerate(row) if c is not None and str(c).strip() != ""]


# ── SRE Summary ──────────────────────────────────────────────────────────────

# metric -> label keys (norm()) it may appear as. The Rental Activity labels
# are the ones on CubeSmart's SRE sheet; the Revenue & AR ones cover the
# common spellings and are checked against a real export before go-live
# (docs/Third-Party-Properties.md). An unknown label is reported, never lost:
# it is kept in the row's `extra`.
SRE_FIELDS: Dict[str, Tuple[str, ...]] = {
    "total_cubes_avail": ("tot cubes avail", "total cubes avail", "total cubes available", "cubes available"),
    "rented_during_month": ("rented during month", "rentals during month", "rented"),
    "vacated_during_month": ("vacated during month", "vacates during month", "vacated"),
    "net_rentals": ("net rentals",),
    "total_company_cubes": ("total company cubes", "tot company cubes", "company cubes"),
    "total_maint_cubes": ("tot maint cubes", "total maint cubes", "total maintenance cubes", "maint cubes"),
    "cubes_occ_eom": ("cubes occ at eom", "cubes occupied at eom", "cubes occ eom", "occupied cubes eom"),
    "pct_rented_cube": ("tot pct rented cube", "total pct rented cube", "pct rented cube"),
    "pct_rented_sqft": ("tot pct rented sqf", "tot pct rented sqft", "tot pct rented sf", "total pct rented sqft", "pct rented sqft"),
    "pct_ncc": ("pct ncc",),
    "pct_cc": ("pct cc",),
    "pct_up": ("pct up",),
    "pct_down": ("pct down",),
    "pct_in": ("pct in",),
    "pct_out": ("pct out",),
    "budget_sf_occupancy": ("budget sf occupancy", "budget sqft occupancy"),
    "los_days": ("los days", "length of stay days"),
    "los_months": ("los months", "length of stay months"),
    "los_value": ("los value", "length of stay value"),
    "rental_revenue": ("rental revenue", "rent revenue", "rental income", "accrued rent", "rent"),
    "fee_revenue": ("fee revenue", "fees", "fee income", "admin fees"),
    "insurance_revenue": ("insurance revenue", "insurance", "protection plan", "tenant protection"),
    "merchandise_revenue": ("merchandise revenue", "merchandise", "retail sales"),
    "other_revenue": ("other revenue", "other income"),
    "total_revenue": ("total revenue", "revenue total", "gross revenue"),
    "promo_discounts": ("promotions", "promo discounts", "promotional discounts", "discounts", "discount", "promo"),
    "promo_move_ins": ("promo move ins", "promotional move ins", "move ins on promo", "promo rentals"),
    "write_offs": ("write offs", "writeoffs", "bad debt", "bad debt write offs"),
    "ar_0_30": ("ar 0 30", "ar 0 30 days", "ar current", "0 30 days", "current ar"),
    "ar_31_60": ("ar 31 60", "ar 31 60 days", "31 60 days"),
    "ar_61_90": ("ar 61 90", "ar 61 90 days", "61 90 days"),
    "ar_91_plus": ("ar 91", "ar 91 plus", "ar 90", "ar 91 days", "91 days", "90 days", "over 90 days", "ar over 90"),
    "ar_total": ("total ar", "ar total", "accounts receivable", "total accounts receivable"),
}
SRE_INT = {"total_cubes_avail", "rented_during_month", "vacated_during_month", "net_rentals", "total_company_cubes",
           "total_maint_cubes", "cubes_occ_eom", "promo_move_ins"}
SRE_PCT = {k for k in SRE_FIELDS if k.startswith("pct_")} | {"budget_sf_occupancy"}
_SRE_BY_LABEL = {a: k for k, al in SRE_FIELDS.items() for a in al}
# Section titles and table furniture: not metrics, not "unknown" either.
_SRE_FURNITURE = {"sre", "rental activity", "revenue ar", "revenue and ar", "values", "month year", "month",
                  "cumulative current customer length of stay", "budget sf occupancy"}


def _sre_value(field: str, v: Any) -> Optional[float]:
    x = num(v, pct=field in SRE_PCT)
    if x is None:
        return None
    return int(round(x)) if field in SRE_INT else x


def parse_sre(sheets) -> Dict[str, Any]:
    """{"stores": {store: {month: {field: value}}}, "unknown": [labels],
    "extras": {store: {month: {label: value}}}}. A store is "" when the export
    does not name one (a single-store export - the upload says which)."""
    stores: Dict[str, Dict[str, Dict[str, Any]]] = {}
    extras: Dict[str, Dict[str, Dict[str, Any]]] = {}
    unknown: List[str] = []
    for _title, rows in sheets:
        store = ""
        colmap: Dict[int, str] = {}
        long_cols: Optional[Dict[int, str]] = None
        month_col = store_col = None
        for row in rows:
            cells = _cells(row)
            if not cells:
                continue
            keys = [norm(c) for _, c in cells]
            # "Store | 6992" or "Store: 6992 - Austin" names the store for what follows
            if keys[0] in ("store", "active store", "store number") and len(cells) >= 2 and store_of(cells[1][1]):
                store = store_of(cells[1][1])
                continue
            months = {i: month_of(c) for i, c in cells}
            months = {i: m for i, m in months.items() if m}
            if len(months) >= 2 and len(months) >= len(cells) - 2:
                colmap = months            # a month header row (wide layout); a later one replaces it
                long_cols = None
                continue
            # long layout header: a Month column and two or more metric columns
            mapped = {i: _SRE_BY_LABEL.get(norm(c)) for i, c in cells}
            if any(norm(c) in ("month", "month year", "period") for _, c in cells) and sum(1 for v in mapped.values() if v) >= 2:
                long_cols = {i: f for i, f in mapped.items() if f}
                month_col = next(i for i, c in cells if norm(c) in ("month", "month year", "period"))
                store_col = next((i for i, c in cells if norm(c) in ("store", "store number", "store id")), None)
                colmap = {}
                continue
            if long_cols is not None:
                m = month_of(row[month_col] if month_col < len(row) else None)
                if not m:
                    continue
                st = store_of(row[store_col]) if store_col is not None and store_col < len(row) else store
                rec = stores.setdefault(st, {}).setdefault(m, {})
                for i, f in long_cols.items():
                    if i < len(row):
                        x = _sre_value(f, row[i])
                        if x is not None:
                            rec[f] = x
                continue
            if not colmap:
                continue
            first_col = min(colmap)
            label_cells = [(i, c) for i, c in cells if i < first_col and isinstance(c, str)]
            if not label_cells:
                continue
            label = label_cells[-1][1]
            key = norm(label)
            field = _SRE_BY_LABEL.get(key)
            values = {colmap[i]: row[i] for i in colmap if i < len(row)}
            if not any(num(v) is not None for v in values.values()):
                continue                    # a section title row ("Revenue & AR") or a blank line
            if field is None:
                if key and key not in _SRE_FURNITURE:
                    if label.strip() not in unknown:
                        unknown.append(label.strip())
                    for m, v in values.items():
                        x = num(v)
                        if x is not None:
                            extras.setdefault(store, {}).setdefault(m, {})[label.strip()[:80]] = x
                continue
            for m, v in values.items():
                x = _sre_value(field, v)
                if x is not None:
                    stores.setdefault(store, {}).setdefault(m, {})[field] = x
    return {"stores": stores, "unknown": unknown, "extras": extras}


# ── Daily Flash ──────────────────────────────────────────────────────────────

_FLASH_COLS = {
    "mtd": "mtd", "budget": "budget", "prior year": "py", "py": "py", "variance": None,
    "daily": "day", "daily budget": "day_budget", "mtd budget": "mtd_budget", "mtd prior year": "mtd_py",
    "mtd py": "mtd_py", "cubes": "cubes", "units": "cubes", "unit pct": None, "sqft": "sqft", "sq ft": "sqft",
    "sqft pct": None, "sq ft pct": None,
}
_MOVE_IN = ("moveins", "move ins", "move in", "rentals")
_MOVE_OUT = ("moveouts", "move outs", "move out", "vacates")
_STATUS = {"occupied": "occupied", "vacant": "vacant", "reserved": "reserved", "damaged": "damaged",
           "unrentable": "unrentable", "totals": "total", "total": "total"}
_REVENUE = {"accrued rent": "accrued_rent", "rent": "accrued_rent", "fees": "fees", "insurance": "insurance",
            "merchandise": "merchandise", "discount": "discount", "discounts": "discount",
            "write offs": "write_offs", "writeoffs": "write_offs", "net revenue": "net_revenue"}


def _flash_field(section: str, label: str, col: str) -> Optional[str]:
    """(table, row label, column) -> a ThirdPartyFlashDay column name."""
    if col in ("day", "day_budget", "mtd_budget", "mtd_py") or (label in _MOVE_IN + _MOVE_OUT and col in ("mtd", "budget", "py")):
        who = "move_ins" if label in _MOVE_IN else "move_outs" if label in _MOVE_OUT else None
        if not who:
            return None
        col = {"budget": "mtd_budget", "py": "mtd_py"}.get(col, col)
        return f"{who}_{col}"
    if col in ("cubes", "sqft"):
        st = _STATUS.get(label)
        return f"{st}_{col}" if st else None
    if col in ("mtd", "budget", "py"):
        if label in _REVENUE:
            return f"{_REVENUE[label]}_{col}"
        if label in ("totals", "total", ""):
            if "allowance" in section:
                return None                 # the allowances total is discount + write-offs, kept separately
            if "net revenue" in section:
                return f"net_revenue_{col}"
            if label and "revenue" in section:
                return f"revenue_{col}"
    return None


# The long layout (one row per day): the column is the field name in words.
def _flat_flash_aliases() -> Dict[str, str]:
    from third_party_models import ThirdPartyFlashDay
    out = {}
    for c in ThirdPartyFlashDay.__table__.columns:
        if c.name in ("id", "property_id", "day", "extra", "import_id", "updated_at"):
            continue
        words = c.name.replace("_", " ")
        out[norm(words)] = c.name
        out[norm(words.replace(" py", " prior year"))] = c.name
        out[norm(words.replace(" day", " daily"))] = c.name
        out[norm(words.replace("move ins", "moveins").replace("move outs", "moveouts"))] = c.name
    return out


def parse_flash(sheets, default_day: Optional[str] = None) -> Dict[str, Any]:
    """{"stores": {store: {day: {field: value}}}, "unknown": [labels]}."""
    flat = _flat_flash_aliases()
    stores: Dict[str, Dict[str, Dict[str, Any]]] = {}
    unknown: List[str] = []
    for _title, rows in sheets:
        store, day, section = "", default_day, ""
        colmap: Dict[int, str] = {}
        long_cols: Optional[Dict[int, str]] = None
        day_col = store_col = None
        for row in rows:
            cells = _cells(row)
            if not cells:
                continue
            keys = [norm(c) for _, c in cells]
            if keys[0] in ("store", "store number") and len(cells) >= 2 and store_of(cells[1][1]):
                store = store_of(cells[1][1])
                continue
            if keys[0] in ("date", "as of", "day", "as of date") and len(cells) >= 2 and day_of(cells[1][1]):
                day = day_of(cells[1][1])
                continue
            # long layout header: a Date column and two or more field columns
            if any(k in ("date", "day", "as of", "as of date") for k in keys) and sum(1 for k in keys if k in flat) >= 2:
                long_cols = {i: flat[norm(c)] for i, c in cells if norm(c) in flat}
                day_col = next(i for i, c in cells if norm(c) in ("date", "day", "as of", "as of date"))
                store_col = next((i for i, c in cells if norm(c) in ("store", "store number", "store id")), None)
                continue
            if long_cols is not None:
                d = day_of(row[day_col] if day_col < len(row) else None)
                if not d:
                    continue
                st = store_of(row[store_col]) if store_col is not None and store_col < len(row) else store
                rec = stores.setdefault(st, {}).setdefault(d, {})
                for i, f in long_cols.items():
                    x = num(row[i]) if i < len(row) else None
                    if x is not None:
                        rec[f] = int(round(x)) if f.endswith(("_cubes", "_sqft")) or f.startswith("move_") else x
                continue
            if len(cells) == 1 and isinstance(cells[0][1], str):
                section = keys[0]
                when = as_of_in_text(cells[0][1])
                if when:
                    day = when
                continue
            cols = {i: _FLASH_COLS.get(k, "?") for (i, _), k in zip(cells, keys)}
            known = {i: c for i, c in cols.items() if c != "?"}
            if len(known) >= 2 and len(known) >= len(cells) - 1:
                colmap = {i: c for i, c in known.items() if c}
                head = [(i, k) for (i, _), k in zip(cells, keys) if cols[i] == "?"]
                if head:
                    section = head[0][1]   # "Revenue", "Allowances", "Occupancy by Status"
                continue
            if not colmap:
                for _, c in cells:
                    when = as_of_in_text(c)
                    if when:
                        day = when
                continue
            first_col = min(colmap)
            label_cells = [(i, c) for i, c in cells if i < first_col and isinstance(c, str)]
            label = norm(label_cells[-1][1]) if label_cells else ""
            if not day:
                raise ParseError("The Daily Flash does not say which day it is for. Pick the As Of date and upload again.")
            hit = False
            for i, col in colmap.items():
                if i >= len(row):
                    continue
                f = _flash_field(section, label, col)
                x = num(row[i])
                if f and x is not None:
                    stores.setdefault(store, {}).setdefault(day, {})[f] = int(round(x)) if f.endswith(("_cubes", "_sqft")) or f.startswith("move_") else x
                    hit = True
            if not hit and label and label not in ("net rentals", "totals", "total") and label_cells:
                if label_cells[-1][1].strip() not in unknown:
                    unknown.append(label_cells[-1][1].strip())
    return {"stores": stores, "unknown": unknown}


# ── Rent roll ────────────────────────────────────────────────────────────────

CUBE = re.compile(r"[0-9A-Z-]{1,12}")
SIZE = re.compile(r"\d+(\.\d+)?\s*[Xx]\s*\d+(\.\d+)?")
_RR_COLS = {
    "store": ("store", "store number", "store id", "site"),
    "cube": ("cube", "unit", "cube number", "unit number", "space"),
    "size": ("cube dimensions", "dimensions", "unit size", "size", "cube size"),
    "attr": ("cube attribute", "attribute", "cube attributes", "unit attribute"),
}


def _rr_header(cells) -> Optional[Dict[str, int]]:
    found = {}
    for i, c in cells:
        k = norm(c)
        for field, names in _RR_COLS.items():
            if k in names and field not in found:
                found[field] = i
    return found if "cube" in found and "size" in found else None


def parse_rent_roll(sheets, store: str = "") -> Dict[str, Any]:
    """{"stores": {store: [{"n", "size", "attr"}]}, "asOf": day or None,
    "bad": count}. Reads four columns found by their headers; with no header
    on a sheet named "Rent Roll", columns A-D (CubeSmart's layout - the same
    reading cubesmart_units.py does). A row whose Cube is not unit-shaped or
    whose size is not WxD is counted as bad and never echoed."""
    out: Dict[str, List[Dict[str, str]]] = {}
    bad = 0
    as_of = None
    seen_header = False
    for title, rows in sheets:
        head = None
        positional = norm(title) == "rent roll"
        for row in rows[:12]:
            cells = _cells(row)
            if head is None:
                head = _rr_header(cells)
            for _, c in cells:
                when = as_of_in_text(c)
                if when and not as_of:
                    as_of = when
        if head is None and not positional:
            continue
        if head is None:
            head = {"store": 0, "cube": 1, "size": 2, "attr": 3}
        else:
            positional = False       # a header was found: data starts below it, never above
        seen_header = True
        started = False
        for row in rows:
            cells = _cells(row)
            if not cells:
                continue
            if not started:
                h = _rr_header(cells)
                if h:
                    head, started = h, True
                    continue
                if not positional:
                    continue
                started = True

            def at(field):
                i = head.get(field)
                return row[i] if i is not None and i < len(row) else None

            st = store_of(at("store")) if "store" in head else ""
            if store and st and st != store:
                continue
            cube, size = at("cube"), at("size")
            if cube is None or str(cube).strip() == "":
                continue
            n = str(cube).strip().upper()
            if isinstance(cube, float) and cube.is_integer():
                n = str(int(cube))
            sz = str(size or "").strip().upper().replace(" ", "")
            if norm(n) in ("total", "totals", "grand total"):
                continue
            if not CUBE.fullmatch(n) or not SIZE.fullmatch(sz):
                bad += 1
                continue
            attr = str(at("attr") or "").strip().upper()[:4]
            out.setdefault(st, []).append({"n": n, "size": sz, "attr": attr if re.fullmatch(r"[A-Z]{0,4}", attr) else ""})
    if not seen_header:
        raise ParseError("No rent roll found: the file needs Cube and Cube Dimensions columns (a Rent Roll sheet).")
    return {"stores": out, "asOf": as_of, "bad": bad}


# ── Which report is it ───────────────────────────────────────────────────────

def detect_kind(sheets) -> Optional[str]:
    for title, rows in sheets:
        if norm(title) == "rent roll":
            return "rent_roll"
        for row in rows[:15]:
            if _rr_header(_cells(row)):
                return "rent_roll"
    sre = parse_sre(sheets)
    if any(sre["stores"].values()):
        return "sre"
    try:
        flash = parse_flash(sheets, default_day="1900-01-01")
    except ParseError:
        return None
    if any(flash["stores"].values()):
        return "flash"
    return None
