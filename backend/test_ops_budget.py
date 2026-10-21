"""Operations > Ops Budget (Oct 21 2026): the calculation rules, the workbook
loader (on a workbook generated here - the real template has not been seen
yet), the API and who may use it. Runs on the throwaway SQLite conftest gives
the run. Run alone: python -m pytest test_ops_budget.py"""
import io
from datetime import date

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

import auth
import database
import models  # noqa: F401 - the shared tables, for create_all
import ops_budget as ob
import ops_budget_workbook as wbk
from database import SessionLocal
from ops_budget_models import OpsBudgetFacility, OpsBudgetLeadMark, OpsBudgetLoad, OpsMetricMonth
from routers import ops_budget as router_mod

TABLES = (OpsMetricMonth, OpsBudgetLoad, OpsBudgetLeadMark, OpsBudgetFacility)
ADMIN = {"email": "admin@example.com", "level": 4, "role": "administrator"}


@pytest.fixture(autouse=True)
def clean():
    models.Base.metadata.create_all(bind=database.engine)
    db = SessionLocal()
    for t in TABLES:
        db.query(t).delete()
    db.commit()
    db.close()
    yield


def _client(user=ADMIN):
    app = FastAPI()
    app.include_router(router_mod.router)
    app.dependency_overrides[auth.get_current_user] = lambda: user
    return TestClient(app)


# ── A generated workbook in the template's shape ─────────────────────────────

SS_HEADERS = ["Month", "Total rentable units", "Occupied units", "Total rentable sq ft", "Occupied sq ft",
              "Gross potential rent $", "Actual rental income $", "Avg standard rate $/mo", "Avg web rate $/mo",
              "Move-ins", "Move-outs", "Sq ft moved in", "Sq ft moved out", "Website visitors", "Leads",
              "Reservations", "Reservations converted to move-ins", "New rentals with insurance / protection",
              "Tenants on autopay", "New online reviews", "Average review rating", "Past due 1-30 days $",
              "Past due 31-60 days $", "Past due 61-90 days $", "Past due 90+ days $", "Liens completed"]


def _month_row(mo, units=712, occ=684, income=151900, ins=26, outs=21, leads=64, res=31, conv=18, insured=19,
               autopay=498, ar=(4820, 1640, 610, 290), liens=2, std=236.5, web=214.8, gpr=168400, sqft=(99850, 96210)):
    return [mo, units, occ, sqft[0], sqft[1], gpr, income, std, web, ins, outs, 3450, 2980, None, leads, res, conv,
            insured, autopay, None, None, *ar, liens]


def make_workbook(name="Greens Storage Escondido", budget_start=date(2026, 1, 1), history_start=date(2023, 1, 1),
                  actual_rows=None, budget_rows=None, ss_headers=SS_HEADERS, vs=True, extra_header=None):
    wb = Workbook()
    setup = wb.active
    setup.title = "Setup"
    setup["A3"], setup["B3"] = "Facility name", name
    setup["A4"], setup["B4"] = "First month of budget year", budget_start
    setup["A5"], setup["B5"] = "First month of history", history_start

    def tab(title, headers, rows):
        ws = wb.create_sheet(title)
        ws["A1"] = f"{title} - inputs are yellow"
        heads = list(headers) + ([extra_header] if extra_header else [])
        for c, h in enumerate(heads, start=1):
            ws.cell(row=5, column=c, value=h)
        for r, row in enumerate(rows, start=6):
            for c, v in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=v)

    tab("SS Actual", ss_headers, actual_rows if actual_rows is not None else [
        _month_row(date(2025, 9, 1), occ=671, income=141200, ins=22, outs=19, leads=58, res=27, conv=14, insured=15,
                   autopay=464, ar=(5310, 2050, 880, 640), liens=3, std=224.9, web=205.1, gpr=159800, sqft=(99850, 93900)),
        _month_row(date(2026, 9, 1)),
    ])
    tab("SS Budget", ss_headers, budget_rows if budget_rows is not None else [
        _month_row(date(2026, m, 1), occ=690, income=152500, ins=24, outs=18, leads=60, res=30, conv=18, insured=18,
                   autopay=511, ar=(4500, 1500, 600, 300), liens=2, std=238.0, web=216.0, sqft=(99850, 96500))
        for m in range(1, 13)])
    if vs:
        vs_heads = ["Month", "Total rentable units", "Occupied units", "Move-ins", "Move-outs", "Tows completed"]
        tab("VS Actual", vs_heads, [[date(2026, 9, 1), 120, 110, 4, 2, 3]])
    wb.create_sheet("SS Report")["A1"] = "=SUM(1,2)"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _upload(client, raw, commit=True, **form):
    data = {"commit": "true" if commit else "false", **form}
    return client.post("/ops-budget/upload", files={"file": ("Ops_Budget_Escondido.xlsx", raw)}, data=data)


# ── Rates and formulas ────────────────────────────────────────────────────────

def test_every_rate_is_recomputed_from_counts():
    v = ob.derive({"total_units": 712, "occupied_units": 684, "total_sqft": 99850, "occupied_sqft": 96210,
                   "gpr": 168400, "rental_income": 151900, "avg_standard_rate": 236.5, "avg_web_rate": 214.8,
                   "move_ins": 26, "move_outs": 21, "sqft_moved_in": 3450, "sqft_moved_out": 2980,
                   "website_visitors": 1000, "reservations": 31, "reservations_converted": 18, "leads": 64,
                   "insured_rentals": 19, "autopay_tenants": 498,
                   "ar_1_30": 4820, "ar_31_60": 1640, "ar_61_90": 610, "ar_90_plus": 290})
    assert round(v["unit_occupancy"] * 100, 1) == 96.1
    assert round(v["sqft_occupancy"] * 100, 1) == 96.4
    assert round(v["economic_occupancy"] * 100, 1) == 90.2
    assert round(v["web_vs_standard"] * 100, 1) == -9.2
    assert v["net_rentals"] == 5 and v["net_sqft"] == 470
    assert round(v["visitor_to_reservation"] * 100, 1) == 3.1
    assert round(v["reservation_to_move_in"] * 100, 1) == 58.1
    assert round(v["lead_to_move_in"] * 100, 1) == 40.6
    assert round(v["insurance_take_rate"] * 100, 1) == 73.1
    assert round(v["autopay_pct"] * 100, 1) == 72.8
    assert v["total_past_due"] == 7360
    assert round(v["share_over_30"] * 100, 1) == 34.5


def test_missing_input_or_zero_denominator_is_a_dash_never_zero():
    v = ob.derive({"total_units": 0, "occupied_units": 0, "leads": None, "move_ins": 3, "reservations": 0,
                   "reservations_converted": 0, "ar_1_30": 10, "ar_31_60": None, "ar_61_90": 0, "ar_90_plus": 0})
    assert v["unit_occupancy"] is None
    assert v["lead_to_move_in"] is None
    assert v["reservation_to_move_in"] is None
    assert v["total_past_due"] is None and v["share_over_30"] is None
    assert v["economic_occupancy"] is None and v["web_vs_standard"] is None
    assert v["net_rentals"] is None          # move-outs missing


def test_a_typed_zero_past_due_total_gives_a_dash_share_not_zero_percent():
    v = ob.derive({"ar_1_30": 0, "ar_31_60": 0, "ar_61_90": 0, "ar_90_plus": 0})
    assert v["total_past_due"] == 0 and v["share_over_30"] is None


# ── Periods ───────────────────────────────────────────────────────────────────

def test_fiscal_ytd_runs_from_the_fiscal_start_and_prior_year_is_minus_12():
    assert ob.period_months("2026-09", "ytd", 1) == [f"2026-{m:02d}" for m in range(1, 10)]
    assert ob.period_months("2026-09", "ytd", 7) == ["2026-07", "2026-08", "2026-09"]
    assert ob.period_months("2026-03", "ytd", 7)[0] == "2025-07"
    assert ob.period_months("2026-09", "month") == ["2026-09"]
    assert ob.add_months("2026-01", -12) == "2025-01" and ob.add_months("2025-12", 1) == "2026-01"


def test_flows_sum_snapshots_take_the_last_month_and_a_blank_month_blanks_a_flow():
    a = ob.aggregate([{"move_ins": 3, "total_units": 100, "leads": 5}, {"move_ins": 4, "total_units": 101, "leads": None}])
    assert a["move_ins"] == 7 and a["total_units"] == 101 and a["leads"] is None


def _rows(fac="escondido", seg="self_storage", **months):
    """months: {"2026_09": {"scenario/source": {metric: value}}}"""
    out = []
    for mk, by in months.items():
        mo = mk.replace("_", "-")
        for ss, vals in by.items():
            scenario, source = ss.split("/")
            for k, v in vals.items():
                out.append({"scenario": scenario, "month": f"{mo}-01", "metric": k, "source": source, "value": v})
    return out


FAC = {"id": "escondido", "shortName": "Escondido", "fiscalStartMonth": 1, "budgetYearStart": "2026-01"}


def _metric(rep, key):
    return next(m for m in rep["metrics"] if m["key"] == key)


def test_ytd_sums_flows_and_budget_shows_only_inside_the_budget_year():
    rows = _rows(**{
        "2026_01": {"actual/workbook": {"move_ins": 10, "total_units": 100, "occupied_units": 90},
                    "budget/workbook": {"move_ins": 12}},
        "2026_02": {"actual/workbook": {"move_ins": 11, "total_units": 100, "occupied_units": 92},
                    "budget/workbook": {"move_ins": 12}},
        "2025_01": {"actual/workbook": {"move_ins": 8, "total_units": 100, "occupied_units": 85}},
        "2025_02": {"actual/workbook": {"move_ins": 9, "total_units": 100, "occupied_units": 86}},
    })
    rep = ob.build_report(ob.table(rows), facility=FAC, segment="self_storage", month="2026-02", period="ytd")
    mi = _metric(rep, "move_ins")
    assert mi["actual"] == 21 and mi["budget"] == 24 and mi["lastYear"] == 17
    occ = _metric(rep, "unit_occupancy")
    assert occ["actual"] == pytest.approx(0.92) and occ["lastYear"] == pytest.approx(0.86)
    # A month outside the budget year has no budget, even if a row exists.
    rep2 = ob.build_report(ob.table(rows), facility={**FAC, "budgetYearStart": "2027-01"}, segment="self_storage",
                           month="2026-02", period="month")
    assert _metric(rep2, "move_ins")["budget"] is None and rep2["budgetAvailable"] is False


# ── Precedence ────────────────────────────────────────────────────────────────

def test_cubby_wins_from_sep_2026_the_workbook_before_and_tows_are_workbook_only():
    assert ob.resolve({"workbook": 5, "cubby_api": 7}, "move_ins", "2026-09") == (7, "cubby_api", 5)
    assert ob.resolve({"workbook": 5, "cubby_api": 7}, "move_ins", "2026-08") == (5, "workbook", 7)
    # Cubby has no value: the workbook's counts.
    assert ob.resolve({"workbook": 5, "cubby_api": None}, "rental_income", "2026-10") == (5, "workbook", None)
    assert ob.resolve({"workbook": 2, "cubby_api": 9}, "tows", "2026-10") == (2, "workbook", None)
    assert ob.resolve({"workbook": 3, "cubby_api": 9}, "move_ins", "2026-10", "budget") == (3, "workbook", None)
    assert ob.resolve({}, "move_ins", "2026-10") == (None, "", None)


def test_the_report_marks_where_an_actual_came_from():
    rows = _rows(**{"2026_09": {"actual/workbook": {"move_ins": 20}, "actual/cubby_api": {"move_ins": 26}}})
    m = _metric(ob.build_report(ob.table(rows), facility=FAC, segment="self_storage", month="2026-09", period="month"),
                "move_ins")
    assert m["actual"] == 26 and m["source"] == "cubby_api" and m["otherSource"] == 20


# ── Variance ──────────────────────────────────────────────────────────────────

def test_variance_direction_and_colors():
    assert ob.variance("move_ins", 26, 24)["status"] == "better"
    v = ob.variance("move_outs", 21, 18)
    assert v["status"] == "worse" and v["arrow"] == "up"                  # more move-outs is worse, arrow still up
    assert ob.variance("ar_90_plus", 290, 300)["status"] == "better"
    assert ob.variance("gpr", 168400, 159800)["status"] == "neutral"      # neutral is never colored
    assert ob.variance("liens", 2, 3) == {**ob.variance("liens", 2, 3), "status": "neutral", "arrow": "down"}
    assert ob.variance("total_units", 712, 712)["status"] == "even"
    v = ob.variance("unit_occupancy", 0.961, 0.969)
    assert v["status"] == "worse" and round(v["pts"], 1) == -0.8 and v["pctChange"] is None
    v = ob.variance("rental_income", 151900, 141200)
    assert v["diff"] == 10700 and round(v["pctChange"] * 100, 1) == 7.6
    assert ob.variance("move_ins", None, 3) is None and ob.variance("move_ins", 3, None) is None
    # Under the shown precision = even.
    assert ob.variance("unit_occupancy", 0.96101, 0.96099)["status"] == "even"


# ── The summary ───────────────────────────────────────────────────────────────

ACTUAL = ob.derive({"total_units": 712, "occupied_units": 684, "gpr": 168400, "rental_income": 151900,
                    "avg_standard_rate": 236.5, "move_ins": 26, "move_outs": 21, "reservations": 31,
                    "reservations_converted": 18, "leads": 64, "insured_rentals": 19, "autopay_tenants": 498,
                    "ar_1_30": 4820, "ar_31_60": 1640, "ar_61_90": 610, "ar_90_plus": 290})
PRIOR = ob.derive({"total_units": 712, "occupied_units": 671, "gpr": 159800, "rental_income": 141200,
                   "avg_standard_rate": 224.9})
BUDGET = ob.derive({"total_units": 712, "occupied_units": 690, "move_ins": 24, "move_outs": 18, "reservations": 30,
                    "reservations_converted": 18, "insured_rentals": 18, "autopay_tenants": 511,
                    "ar_1_30": 4500, "ar_31_60": 1500, "ar_61_90": 600, "ar_90_plus": 300})


def test_summary_headline_and_lists():
    s = ob.summary("self_storage", "month", "2026-09", ACTUAL, BUDGET, PRIOR)
    assert s["headline"] == ("Self-storage is 96.1% full and earning $10,700 more rent than the same month last "
                             f"year, but {s['attentionCount']} items need attention.")
    good = [i["text"] for i in s["goingWell"]]
    assert good[0] == "Unit occupancy is 96.1%, up 1.8 pts from last year."
    assert "Rental income is $151,900, up $10,700 from last year." in good
    assert "Standard rate is $236.50, up $11.60 from last year." in good
    bad = [i["text"] for i in s["needsAttention"]]
    assert bad[:2] == ["Net rentals is 5 against a plan of 6.", "Move-outs is 21 against a plan of 18."]
    assert "Reservation conversion is 58.1% against a plan of 60.0%." in bad
    assert len(s["needsAttention"]) <= 4 and len(s["goingWell"]) <= 4
    assert "Economic occupancy is 90.2%, up 1.8 pts from last year." in good
    # each metric once
    keys = [i["metric"] for i in s["goingWell"] + s["needsAttention"]]
    assert len(keys) == len(set(keys))


def test_summary_ytd_wording_less_rent_and_one_item():
    a = ob.derive({"total_units": 100, "occupied_units": 90, "rental_income": 900, "move_outs": 5})
    p = ob.derive({"rental_income": 1000})
    b = ob.derive({})
    s = ob.summary("vehicle", "ytd", "2026-09", a, b, p)
    assert [i["metric"] for i in s["needsAttention"]] == ["rental_income"]
    assert s["headline"] == "Vehicle storage is 90.0% full and earning $100 less rent than this time last year, but 1 item needs attention."
    s = ob.summary("vehicle", "month", "2026-09", ob.derive({}), ob.derive({}), ob.derive({}))
    assert s["headline"] == "There are not enough figures for vehicle storage in September 2026 yet."
    assert s["goingWell"] == [] and s["needsAttention"] == []


def test_takeaways():
    assert ob.funnel_takeaway({"leads": 64, "move_ins": 26}).startswith("Out of every 10 inquiries, about 4 became a rental.")
    assert ob.funnel_takeaway({"leads": None, "move_ins": 26}) == ""
    pts = [{"move_ins": 22, "move_outs": 20}] * 12
    assert ob.inout_takeaway(pts) == "264 customers moved in and 240 moved out over the last 12 months, a net gain of 24."
    assert ob.inout_takeaway([{"move_ins": 1, "move_outs": 3}] * 12).endswith("a net loss of 24.")
    occ = [{"unit_occupancy": 0.924}] + [{"unit_occupancy": None}] * 22 + [{"unit_occupancy": 0.961}]
    assert ob.occupancy_takeaway(occ).startswith("Units rented rose from 92.4% to 96.1% over two years. The gap")
    rates = [{"avg_standard_rate": 224.8, "avg_web_rate": 200}] * 12 + [{"avg_standard_rate": 236.5, "avg_web_rate": 214.8}]
    assert ob.rate_takeaway(rates) == "The standard rate is up 5.2% from a year ago. Online customers pay about $22 a month less."


# ── The workbook loader ───────────────────────────────────────────────────────

def test_headers_match_tolerantly_and_unknown_or_missing_are_reported():
    assert wbk.norm("Avg  Standard Rate ($/mo)") == wbk.norm("avg standard rate")
    assert wbk.INPUT_HEADERS[wbk.norm("PAST DUE 90+ DAYS $")] == "ar_90_plus"
    assert wbk.INPUT_HEADERS[wbk.norm("New rentals with insurance/protection")] == "insured_rentals"
    raw = make_workbook(extra_header="Mystery Column")
    db = SessionLocal()
    try:
        p = wbk.build_preview(db, raw, "x.xlsx")
    finally:
        db.close()
    assert p["errors"] == [] and p["facility"]["id"] == "escondido"
    ss = next(s for s in p["sheets"] if s["sheet"] == "SS Actual")
    assert "Mystery Column" in ss["unknown"] and "Mystery Column" in p["unknownHeaders"]
    assert ss["missing"] == []
    vs = next(s for s in p["sheets"] if s["sheet"] == "VS Actual")
    assert "Leads" in vs["missing"] and "Tows completed" in vs["columns"]
    assert "There is no VS Budget tab, so nothing was read for it." in p["warnings"]


def test_a_calculated_rate_column_is_ignored_and_said_so():
    heads = SS_HEADERS[:3] + ["Unit occupancy %"]
    raw = make_workbook(ss_headers=heads, actual_rows=[[date(2026, 9, 1), 100, 90, 0.9]], budget_rows=[])
    db = SessionLocal()
    try:
        p = wbk.build_preview(db, raw, "x.xlsx")
    finally:
        db.close()
    ss = next(s for s in p["sheets"] if s["sheet"] == "SS Actual")
    assert ss["ignored"] == ["Unit occupancy %"] and ss["unknown"] == []


def test_blank_is_null_typed_zero_is_zero_and_reupload_upserts():
    c = _client()
    rows = [_month_row(date(2026, 9, 1), liens=0)]
    rows[0][SS_HEADERS.index("Leads")] = None
    r = _upload(c, make_workbook(actual_rows=rows, vs=False))
    assert r.status_code == 200, r.text
    assert r.json()["saved"] is True
    db = SessionLocal()
    try:
        def val(metric, month="2026-09-01"):
            row = db.query(OpsMetricMonth).filter_by(facility_id="escondido", segment="self_storage", month=month,
                                                     scenario="actual", metric=metric, source="workbook").first()
            return "absent" if row is None else row.value
        assert val("liens") == 0.0
        assert val("leads") == "absent"
        assert val("move_ins") == 26
        fac = db.get(OpsBudgetFacility, "escondido")
        assert fac.budget_year_start == "2026-01" and fac.fiscal_start_month == 1 and fac.history_start == "2023-01"
        n = db.query(OpsMetricMonth).count()
        db.expire_all()
    finally:
        db.close()
    # Re-upload: move-ins changed, liens now blank -> the row is cleared to NULL, nothing duplicated.
    rows = [_month_row(date(2026, 9, 1), ins=30, liens=None)]
    rows[0][SS_HEADERS.index("Leads")] = None
    r = _upload(c, make_workbook(actual_rows=rows, vs=False))
    assert r.status_code == 200, r.text
    db = SessionLocal()
    try:
        q = db.query(OpsMetricMonth).filter_by(facility_id="escondido", segment="self_storage", month="2026-09-01",
                                               scenario="actual", source="workbook")
        assert q.filter_by(metric="move_ins").one().value == 30
        assert q.filter_by(metric="liens").one().value is None
        assert db.query(OpsMetricMonth).count() == n
        loads = db.query(OpsBudgetLoad).filter_by(kind="workbook").all()
        assert len(loads) == 2 and all(lo.file_sha256 and lo.file_size and lo.loaded_by == ADMIN["email"] for lo in loads)
    finally:
        db.close()


def test_preview_saves_nothing_and_text_in_a_number_cell_is_reported():
    c = _client()
    rows = [_month_row(date(2026, 9, 1))]
    rows[0][SS_HEADERS.index("Move-ins")] = "lots"
    r = _upload(c, make_workbook(actual_rows=rows, vs=False), commit=False)
    assert r.status_code == 200 and r.json()["saved"] is False
    assert any("'lots' is" in w for w in r.json()["warnings"])
    assert "values" not in r.json()
    db = SessionLocal()
    try:
        assert db.query(OpsMetricMonth).count() == 0
    finally:
        db.close()


def test_months_read_in_any_common_form_and_budget_rows_outside_the_year_skip():
    assert wbk.month_of("Sep 2026") == "2026-09" and wbk.month_of("2026-09") == "2026-09"
    assert wbk.month_of("9/1/2026") == "2026-09" and wbk.month_of(46266) == "2026-09"
    assert wbk.month_of("Total") is None
    budget = [_month_row(date(2025, 12, 1)), _month_row(date(2026, 1, 1))]
    db = SessionLocal()
    try:
        p = wbk.build_preview(db, make_workbook(budget_rows=budget, vs=False), "x.xlsx")
    finally:
        db.close()
    sb = next(s for s in p["sheets"] if s["sheet"] == "SS Budget")
    assert sb["months"] == 1 and any("outside the budget year" in w for w in sb["warnings"])


def test_an_unknown_site_or_a_wrong_pick_is_refused():
    c = _client()
    r = _upload(c, make_workbook(name="Greens Storage Austin"))
    assert r.status_code == 400 and "could not tell which site" in r.json()["detail"]
    r = _upload(c, make_workbook(), facility_id="temecula")
    assert r.status_code == 400 and "not Temecula" in r.json()["detail"]
    r = c.post("/ops-budget/upload", files={"file": ("x.xlsx", b"not a workbook")}, data={"commit": "true"})
    assert r.status_code == 400 and "Excel" in r.json()["detail"]


# ── The API ───────────────────────────────────────────────────────────────────

def test_report_and_trend_after_an_upload():
    c = _client()
    assert _upload(c, make_workbook()).status_code == 200
    r = c.get("/ops-budget", params={"facility": "escondido", "segment": "self_storage", "month": "2026-09"})
    assert r.status_code == 200, r.text
    body = r.json()
    occ = _metric(body, "unit_occupancy")
    assert round(occ["actual"] * 100, 1) == 96.1 and round(occ["budget"] * 100, 1) == 96.9
    assert round(occ["lastYear"] * 100, 1) == 94.2 and occ["vsLastYear"]["status"] == "better"
    assert body["summary"]["headline"].startswith("Self-storage is 96.1% full and earning $10,700 more rent")
    visitors = _metric(body, "website_visitors")
    assert visitors["pending"] is True and visitors["actual"] is None
    assert not any(m["key"] == "tows" for m in body["metrics"])           # tows: vehicle only
    assert body["budgetAvailable"] is True and body["monthLabel"] == "Sep 2026"

    v = c.get("/ops-budget", params={"facility": "escondido", "segment": "vehicle", "month": "2026-09"}).json()
    assert _metric(v, "tows")["actual"] == 3

    t = c.get("/ops-budget/trend", params={"facility": "escondido", "segment": "self_storage", "month": "2026-09"})
    assert t.status_code == 200
    pts = t.json()["points"]
    assert len(pts) == 24 and pts[-1]["month"] == "2026-09" and round(pts[-1]["unit_occupancy"] * 100, 1) == 96.1

    f = c.get("/ops-budget/facilities").json()
    assert [x["shortName"] for x in f["facilities"]] == ["Escondido", "Fairfield", "Murrieta", "Temecula", "Valley Center"]
    assert f["canUpload"] is True and f["facilities"][0]["dataRange"]["self_storage"]["last"] == "2026-09"

    loads = c.get("/ops-budget/loads").json()["loads"]
    assert loads and loads[0]["fileName"] == "Ops_Budget_Escondido.xlsx"
    got = c.get(f"/ops-budget/loads/{loads[0]['id']}/file")
    assert got.status_code == 200 and got.content[:2] == b"PK"


def test_bad_parameters_are_422_and_unknown_site_404():
    c = _client()
    assert c.get("/ops-budget", params={"facility": "escondido", "segment": "boats"}).status_code == 422
    assert c.get("/ops-budget", params={"facility": "escondido", "period": "week"}).status_code == 422
    assert c.get("/ops-budget", params={"facility": "escondido", "month": "2022-12"}).status_code == 422
    assert c.get("/ops-budget", params={"facility": "nowhere", "month": "2026-09"}).status_code == 404


def test_permissions():
    employee = {"email": "emp@example.com", "level": 1, "role": "employee"}
    supervisor = {"email": "sup@example.com", "level": 2, "role": "supervisor"}
    manager = {"email": "mgr@example.com", "level": 3, "role": "manager"}
    raw = make_workbook()
    assert _client(employee).get("/ops-budget/facilities").status_code == 401
    assert _client(employee).get("/ops-budget", params={"facility": "escondido"}).status_code == 401
    sup = _client(supervisor)
    assert sup.get("/ops-budget/facilities").status_code == 200
    assert sup.get("/ops-budget/facilities").json()["canUpload"] is False
    assert sup.get("/ops-budget", params={"facility": "escondido", "month": "2026-09"}).status_code == 200
    mgr = _client(manager)
    assert _upload(mgr, raw).status_code == 401
    assert mgr.get("/ops-budget/loads").status_code == 401
    assert mgr.post("/ops-budget/snapshot").status_code == 401
    assert _upload(_client(ADMIN), raw).status_code == 200


def test_no_tenant_detail_in_the_tables():
    cols = {c.name for t in TABLES for c in t.__table__.columns}
    assert not cols & {"tenant", "customer_name", "email", "phone", "address", "customer_id", "unit_name"}
