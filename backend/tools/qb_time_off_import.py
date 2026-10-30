"""Seed leave-balance opening rows from a QuickBooks Time Off Report (.xlsx).

QB's Time Off Report lays out one check-row per employee per period, with a
running Vac/Sick Balance. The last data row for each person is their balance
as of the end of the report's date range - that is what Nexus's leave ledger
needs as the opening row.

Usage (from the repo root, with `backend/` on PYTHONPATH):
    python backend/tools/qb_time_off_import.py --company MCD \
        --file "MCD - Time Off Report.xlsx"
    python backend/tools/qb_time_off_import.py --company OSM \
        --file "OSM - Time Off Report.xlsx" --commit

Dry-run by default. The script:
  * parses the workbook (openpyxl)
  * reads backend/tools/qb_time_off_mapping.json to turn QB code|name into a
    Nexus work email - fill the blanks there if HR has not added a person
    yet, or set skip: true to leave them out
  * sets the company's cutover (from the mapping file's `cutover` field) if
    missing - never changes an existing one
  * calls leave_balances.set_opening with each matched person's final Vac
    and Sick balance (hours)
  * prints a per-person report: WRITE / SKIP (reason)

Nothing is written unless `--commit` is passed.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
BACKEND = HERE.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

def _lazy_db():
    """Import DB-only pieces on demand - lets --parse-only run with no DB."""
    import leave_balances as lb  # noqa: E402
    import models  # noqa: E402
    from database import SessionLocal  # noqa: E402
    return lb, models, SessionLocal


HR_ENTITY_IDS = {
    "MCD": "7fcbe8bb-63af-4c5d-9696-1f2c8461b154",
    "OSM": "52eaa7a3-c3ff-4072-a9f4-5d3c4072279e",
}


def _parse_workbook(path: Path) -> list[dict]:
    """Return [{code, name, vac, sick}, ...] - one entry per employee, carrying
    the LAST per-check Vac Balance / Sick Balance (col H / col K)."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True)
    ws = wb["Time Off"]
    out: list[dict] = []
    cur: dict | None = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        # Banner row: a name in A, nothing in B. "Date" is the header, skip it.
        if isinstance(a, str) and a != "Date" and b is None:
            if cur:
                out.append(cur)
            m = re.match(r"^(.*?)\s*\(+\s*(E-\d+)\s*\)+\s*$", a.strip())
            name, code = (m.group(1).strip(), m.group(2)) if m else (a.strip(), "")
            # Strip stray leftover punctuation and middle initials from the banner.
            name = re.sub(r"\s*[(\[]+\s*$", "", name).strip()
            name = re.sub(r"\s+[A-Z]\.?\s+", " ", name).strip()
            cur = {"code": code, "name": name, "vac": 0.0, "sick": 0.0}
            continue
        if isinstance(a, str) and a == "Date":
            continue
        if a is None and b is None:
            # Totals row - sums, not the ending balance.
            continue
        if hasattr(a, "year") and cur is not None:
            cur["vac"] = float(ws.cell(r, 8).value or 0)
            cur["sick"] = float(ws.cell(r, 11).value or 0)
    if cur:
        out.append(cur)
    return out


def _load_mapping() -> dict:
    with open(HERE / "qb_time_off_mapping.json") as f:
        return json.load(f)


def _mapping_key(code: str, name: str) -> str:
    return f"{code or ''}|{name}"


def _employee_by_email(db, models, email: str):
    email = (email or "").strip().lower()
    if not email:
        return None
    return (
        db.query(models.NexusEmployee)
        .filter(models.NexusEmployee.work_email == email)
        .filter((models.NexusEmployee.deleted_at == "") | (models.NexusEmployee.deleted_at.is_(None)))
        .first()
    )


def parse_only(company: str, xlsx: Path) -> int:
    cfg = _load_mapping()
    comp = company.upper()
    rows = cfg.get(comp) or {}
    people = _parse_workbook(xlsx)
    print(f"=== {comp} - parse only (no DB) - {len(people)} people from {xlsx.name} ===\n")
    mapped: list[str] = []
    missing_email: list[str] = []
    not_in_mapping: list[str] = []
    skipped: list[str] = []
    for p in people:
        key = _mapping_key(p["code"], p["name"])
        vac = round(max(0.0, p["vac"]), 4)
        sick = round(max(0.0, p["sick"]), 4)
        label = f"{p['code'] or '-':<8} {p['name']:<30} vac={vac:>8}  sick={sick:>6}"
        cfg_row = rows.get(key)
        if cfg_row is None:
            not_in_mapping.append(label)
            continue
        if cfg_row.get("skip"):
            skipped.append(f"{label}  :: skip")
            continue
        em = (cfg_row.get("email") or "").strip()
        if em:
            mapped.append(f"{label}  -> {em}")
        else:
            missing_email.append(label)

    print(f"MAPPED ({len(mapped)}): would import opening balances for these")
    for m in mapped:
        print(f"  {m}")
    print()
    print(f"NEEDS EMAIL ({len(missing_email)}): edit qb_time_off_mapping.json to add their Nexus work_email")
    for m in missing_email:
        print(f"  {m}")
    print()
    if skipped:
        print(f"SKIPPED ({len(skipped)}):")
        for s in skipped:
            print(f"  {s}")
        print()
    if not_in_mapping:
        print(f"NOT IN MAPPING ({len(not_in_mapping)}) - new names since the mapping was built:")
        for n in not_in_mapping:
            print(f"  {n}")
        print()
    return 0


def run(company: str, xlsx: Path, commit: bool) -> int:
    lb, models, SessionLocal = _lazy_db()
    cfg = _load_mapping()
    cutover = str(cfg.get("cutover") or "").strip()
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", cutover):
        print("ERROR: mapping file is missing a valid `cutover` (YYYY-MM-DD).")
        return 2

    comp = company.upper()
    company_id = HR_ENTITY_IDS.get(comp)
    if not company_id:
        print(f"ERROR: unknown company {comp!r}. Known: {sorted(HR_ENTITY_IDS)}")
        return 2

    rows = cfg.get(comp) or {}
    if not rows:
        print(f"ERROR: mapping file has no entries for {comp}.")
        return 2

    print(f"=== {comp} ({company_id}) - cutover {cutover} ===")
    people = _parse_workbook(xlsx)
    print(f"workbook: {len(people)} people parsed from {xlsx.name}\n")

    db = SessionLocal()
    actor = "qb_time_off_import"
    try:
        existing_cutover = lb.cutover(db, company_id)
        if not existing_cutover:
            print(f"CUTOVER: {comp} has no cutover set - would set to {cutover}")
            if commit:
                lb.set_cutover(db, cutover, actor, company_id)
                db.commit()
                print(f"  SET cutover for {comp} = {cutover}")
        elif existing_cutover != cutover:
            print(
                f"CUTOVER: {comp} already has cutover {existing_cutover} - "
                f"leaving it (mapping says {cutover}). Not changing an existing cutover."
            )
        else:
            print(f"CUTOVER: {comp} cutover already {cutover}")
        print()

        written = 0
        skipped: list[tuple[str, str]] = []
        for p in people:
            key = _mapping_key(p["code"], p["name"])
            cfg_row = rows.get(key)
            vac = round(max(0.0, p["vac"]), 4)
            sick = round(max(0.0, p["sick"]), 4)
            label = f"{p['code'] or '-':<8} {p['name']:<28} vac={vac:>8}  sick={sick:>6}"

            if cfg_row is None:
                skipped.append((label, "no mapping entry (edit qb_time_off_mapping.json)"))
                continue
            if cfg_row.get("skip"):
                skipped.append((label, "skipped by mapping"))
                continue
            email = (cfg_row.get("email") or "").strip().lower()
            if not email:
                skipped.append((label, "email blank in mapping"))
                continue

            emp = _employee_by_email(db, models, email)
            if emp is None:
                skipped.append((label, f"no Nexus employee with work_email {email}"))
                continue
            if (emp.company or "") != company_id:
                # Importing anyway because the email is explicit, but note it.
                note = f"WRITE (NOTE: Nexus has {emp.first_name} {emp.last_name} under company {emp.company or '-'}, not {comp})"
            else:
                note = "WRITE"

            print(f"  {note:<110}  {label}  -> {email}")
            if commit:
                try:
                    lb.set_opening(db, email, vac if vac else None, sick if sick else None, actor)
                    written += 1
                except Exception as e:
                    print(f"    ! {email}: {e}")
                    db.rollback()
                    continue

        if commit:
            db.commit()

        print()
        print(f"SKIPPED ({len(skipped)}):")
        for label, reason in skipped:
            print(f"  - {label}  :: {reason}")
        print()
        print(f"Summary: {'WROTE' if commit else 'WOULD WRITE'} {written if commit else sum(1 for p in people if (rows.get(_mapping_key(p['code'], p['name'])) or {}).get('email'))} people, SKIPPED {len(skipped)}.")
        if not commit:
            print("DRY RUN - no rows changed. Re-run with --commit to apply.")
    finally:
        db.close()
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--company", required=True, choices=sorted(HR_ENTITY_IDS))
    ap.add_argument("--file", required=True, help="Path to the QB Time Off Report .xlsx")
    ap.add_argument("--commit", action="store_true", help="Actually write (default is dry-run)")
    ap.add_argument("--parse-only", action="store_true",
                    help="Report which rows map / do not, without opening a DB connection")
    args = ap.parse_args()
    xlsx = Path(args.file)
    if not xlsx.exists():
        print(f"ERROR: file not found: {xlsx}")
        return 2
    if args.parse_only:
        return parse_only(args.company, xlsx)
    return run(args.company, xlsx, args.commit)


if __name__ == "__main__":
    sys.exit(main())
